import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.http import HttpResponse
from django.utils import timezone
from .models import ExportLog


def export_review_sheet(projects, user, filter_params=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '伦理审查表'

    header_font = Font(bold=True, size=12)
    header_alignment = Alignment(horizontal='center', vertical='center')
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    headers = [
        '课题编号', '课题名称', '负责人', '部门', '状态',
        '提交时间', '补件次数', '每次补件时间', '归档时间',
        '评审意见汇总'
    ]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border

    for row, project in enumerate(projects, 2):
        resubmissions = project.resubmissions.all()
        resubmission_times = '; '.join(
            [r.submitted_at.strftime('%Y-%m-%d %H:%M') for r in resubmissions]
        ) if resubmissions else '-'

        comments = project.review_comments.select_related('clause', 'reviewer')
        comments_by_clause = {}
        for comment in comments:
            clause_key = comment.clause.clause_number if comment.clause else '其他'
            if clause_key not in comments_by_clause:
                comments_by_clause[clause_key] = []
            comments_by_clause[clause_key].append(
                f'[{comment.reviewer.get_full_name()}] {comment.content}'
            )
        comments_summary = '\n'.join(
            f'条款 {k}: {"；".join(v)}' for k, v in comments_by_clause.items()
        ) if comments_by_clause else '-'

        data = [
            project.project_code,
            project.title,
            project.principal_investigator.get_full_name(),
            project.department,
            project.get_status_display(),
            project.submitted_at.strftime('%Y-%m-%d %H:%M') if project.submitted_at else '-',
            resubmissions.count(),
            resubmission_times,
            project.archived_at.strftime('%Y-%m-%d %H:%M') if project.archived_at else '-',
            comments_summary,
        ]

        for col, value in enumerate(data, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='top', wrap_text=True)

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 20
    ws.row_dimensions[1].height = 30

    file_name = f'伦理审查表_{timezone.now().strftime("%Y%m%d_%H%M%S")}.xlsx'

    if filter_params:
        ExportLog.objects.create(
            user=user,
            export_type='review_sheet',
            filter_params=filter_params,
            file_name=file_name,
        )

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{file_name}"'
    wb.save(response)
    return response


def export_project_detail(project, user):
    wb = openpyxl.Workbook()

    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    header_font = Font(bold=True)

    ws1 = wb.active
    ws1.title = '课题信息'

    ws1['A1'] = '课题基本信息'
    ws1['A1'].font = Font(bold=True, size=14)
    ws1.merge_cells('A1:B1')

    info_fields = [
        ('课题编号', project.project_code),
        ('课题名称', project.title),
        ('负责人', project.principal_investigator.get_full_name()),
        ('联系电话', project.principal_investigator.phone),
        ('部门', project.department),
        ('状态', project.get_status_display()),
        ('创建时间', project.created_at.strftime('%Y-%m-%d %H:%M')),
        ('提交时间', project.submitted_at.strftime('%Y-%m-%d %H:%M') if project.submitted_at else '-'),
        ('归档时间', project.archived_at.strftime('%Y-%m-%d %H:%M') if project.archived_at else '-'),
        ('课题简介', project.description),
    ]

    for row, (label, value) in enumerate(info_fields, 3):
        ws1.cell(row=row, column=1, value=label).font = header_font
        ws1.cell(row=row, column=1).border = thin_border
        ws1.cell(row=row, column=2, value=value).border = thin_border

    ws1.column_dimensions['A'].width = 20
    ws1.column_dimensions['B'].width = 60

    ws2 = wb.create_sheet('材料清单')
    ws2['A1'] = '材料清单'
    ws2['A1'].font = Font(bold=True, size=14)

    material_headers = ['材料类型', '当前版本', '最新版本号', '上传时间', '上传者']
    for col, header in enumerate(material_headers, 1):
        cell = ws2.cell(row=3, column=col, value=header)
        cell.font = header_font
        cell.border = thin_border

    materials = project.materials.select_related('material_type', 'current_version').prefetch_related('versions')
    for row, material in enumerate(materials, 4):
        latest = material.get_latest_version()
        data = [
            material.material_type.name,
            f'v{material.current_version.version_number}' if material.current_version else '-',
            f'v{latest.version_number}' if latest else '-',
            latest.created_at.strftime('%Y-%m-%d %H:%M') if latest else '-',
            latest.uploader.get_full_name() if latest else '-',
        ]
        for col, value in enumerate(data, 1):
            ws2.cell(row=row, column=col, value=value).border = thin_border

    for col in range(1, 6):
        ws2.column_dimensions[get_column_letter(col)].width = 20

    ws3 = wb.create_sheet('评审意见')
    ws3['A1'] = '评审意见汇总'
    ws3['A1'].font = Font(bold=True, size=14)

    comment_headers = ['条款', '评审人', '意见内容', '状态', '创建时间']
    for col, header in enumerate(comment_headers, 1):
        cell = ws3.cell(row=3, column=col, value=header)
        cell.font = header_font
        cell.border = thin_border

    comments = project.review_comments.select_related('clause', 'reviewer').order_by('clause__clause_number')
    for row, comment in enumerate(comments, 4):
        data = [
            f'{comment.clause.clause_number} {comment.clause.title}' if comment.clause else '其他',
            comment.reviewer.get_full_name(),
            comment.content,
            comment.get_status_display(),
            comment.created_at.strftime('%Y-%m-%d %H:%M'),
        ]
        for col, value in enumerate(data, 1):
            ws3.cell(row=row, column=col, value=value).border = thin_border

    ws3.column_dimensions['A'].width = 30
    ws3.column_dimensions['B'].width = 15
    ws3.column_dimensions['C'].width = 60
    ws3.column_dimensions['D'].width = 12
    ws3.column_dimensions['E'].width = 18

    ws4 = wb.create_sheet('补件记录')
    ws4['A1'] = '补件记录'
    ws4['A1'].font = Font(bold=True, size=14)

    resub_headers = ['补件时间', '提交者', '回应说明', '回应的意见数', '关联材料版本']
    for col, header in enumerate(resub_headers, 1):
        cell = ws4.cell(row=3, column=col, value=header)
        cell.font = header_font
        cell.border = thin_border

    resubmissions = project.resubmissions.select_related('submitter', 'material_version')
    for row, resub in enumerate(resubmissions, 4):
        data = [
            resub.submitted_at.strftime('%Y-%m-%d %H:%M'),
            resub.submitter.get_full_name(),
            resub.response_note,
            resub.addressed_comments.count(),
            f'{resub.material_version.material.material_type.name} v{resub.material_version.version_number}',
        ]
        for col, value in enumerate(data, 1):
            ws4.cell(row=row, column=col, value=value).border = thin_border

    ws4.column_dimensions['A'].width = 18
    ws4.column_dimensions['B'].width = 15
    ws4.column_dimensions['C'].width = 50
    ws4.column_dimensions['D'].width = 15
    ws4.column_dimensions['E'].width = 30

    file_name = f'课题详情_{project.project_code}_{timezone.now().strftime("%Y%m%d_%H%M%S")}.xlsx'

    ExportLog.objects.create(
        user=user,
        export_type='project_detail',
        filter_params={'project_id': project.id},
        file_name=file_name,
    )

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{file_name}"'
    wb.save(response)
    return response
