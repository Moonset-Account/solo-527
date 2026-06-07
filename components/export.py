import pandas as pd
import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from utils.helpers import get_caliber_note, format_duration


def generate_excel_report(filter_state: dict, kpi_data: dict, pareto_data: dict, 
                          line_data: dict, efficiency_data: dict, correlation_data: dict):
    output = io.BytesIO()
    wb = Workbook()
    
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    ws_summary = wb.active
    ws_summary.title = "报告概览"
    
    ws_summary["A1"] = "工厂设备停机分析报告"
    ws_summary["A1"].font = Font(bold=True, size=16, color="1E3A5F")
    ws_summary.merge_cells("A1:F1")
    
    ws_summary["A3"] = "生成时间"
    ws_summary["B3"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    
    ws_summary["A4"] = "筛选条件"
    filter_desc = []
    if filter_state.get("start_date") and filter_state.get("end_date"):
        filter_desc.append(f"时间范围: {filter_state['start_date']} 至 {filter_state['end_date']}")
    if filter_state.get("line_ids"):
        filter_desc.append(f"产线: {filter_state['line_ids']}")
    if filter_state.get("equipment_ids"):
        filter_desc.append(f"设备: {filter_state['equipment_ids']}")
    if filter_state.get("shift_ids"):
        filter_desc.append(f"班次: {filter_state['shift_ids']}")
    if filter_state.get("breakdown_type"):
        type_map = {"all": "全部", "planned": "仅计划检修", "unplanned": "仅突发故障"}
        filter_desc.append(f"停机类型: {type_map.get(filter_state['breakdown_type'], '全部')}")
    ws_summary["B4"] = "; ".join(filter_desc) if filter_desc else "无筛选(全部数据)"
    
    ws_summary["A6"] = "核心KPI指标"
    ws_summary["A6"].font = Font(bold=True, size=12)
    ws_summary.merge_cells("A6:F6")
    
    kpi_headers = ["指标名称", "数值", "单位", "口径说明"]
    for col, header in enumerate(kpi_headers, 1):
        cell = ws_summary.cell(row=7, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border
    
    kpi_rows = [
        ["总停机时长", format_duration(kpi_data.get("total_duration", 0)), "", get_caliber_note("total_duration")],
        ["停机总次数", kpi_data.get("total_count", 0), "次", ""],
        ["突发停机时长", format_duration(kpi_data.get("unplanned_duration", 0)), "", get_caliber_note("unplanned_duration")],
        ["突发停机占比", f"{kpi_data.get('unplanned_ratio', 0)}%", "", get_caliber_note("breakdown_type")],
        ["平均修复时间(MTTR)", kpi_data.get("mttr", 0), "分钟", get_caliber_note("mttr")],
        ["平均故障间隔(MTBF)", kpi_data.get("mtbf", 0), "小时", get_caliber_note("mtbf")],
        ["设备可用率", f"{kpi_data.get('availability', 0)}%", "", get_caliber_note("availability")],
    ]
    
    for row_idx, row_data in enumerate(kpi_rows, 8):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws_summary.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(wrap_text=True)
    
    ws_summary.column_dimensions["A"].width = 25
    ws_summary.column_dimensions["B"].width = 20
    ws_summary.column_dimensions["C"].width = 15
    ws_summary.column_dimensions["D"].width = 60
    
    if pareto_data.get("items"):
        ws_pareto = wb.create_sheet("Pareto分析")
        ws_pareto["A1"] = "停机原因Pareto分析（按时长排序）"
        ws_pareto["A1"].font = Font(bold=True, size=14, color="1E3A5F")
        ws_pareto.merge_cells("A1:H1")
        
        pareto_headers = ["故障类型", "停机时长(分钟)", "停机次数", "占比(%)", "累积占比(%)", "故障描述", "改善建议"]
        for col, header in enumerate(pareto_headers, 1):
            cell = ws_pareto.cell(row=3, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border
        
        for row_idx, item in enumerate(pareto_data["items"], 4):
            ws_pareto.cell(row=row_idx, column=1, value=item["fault_name"]).border = thin_border
            ws_pareto.cell(row=row_idx, column=2, value=item["duration"]).border = thin_border
            ws_pareto.cell(row=row_idx, column=3, value=item["count"]).border = thin_border
            ws_pareto.cell(row=row_idx, column=4, value=item["percentage"]).border = thin_border
            ws_pareto.cell(row=row_idx, column=5, value=item["cumulative_percentage"]).border = thin_border
            ws_pareto.cell(row=row_idx, column=6, value=item["description"]).border = thin_border
            ws_pareto.cell(row=row_idx, column=7, value=item["suggestion"]).border = thin_border
        
        for col in ["A", "B", "C", "D", "E", "F", "G"]:
            ws_pareto.column_dimensions[col].width = 20
        ws_pareto.column_dimensions["F"].width = 40
        ws_pareto.column_dimensions["G"].width = 40
    
    if line_data.get("lines"):
        ws_line = wb.create_sheet("产线对比")
        ws_line["A1"] = "各产线停机对比"
        ws_line["A1"].font = Font(bold=True, size=14, color="1E3A5F")
        ws_line.merge_cells("A1:H1")
        
        line_headers = ["产线名称", "总停机时长(分钟)", "停机次数", "平均时长(分钟)", 
                        "计划检修时长", "突发故障时长", "突发故障占比(%)"]
        for col, header in enumerate(line_headers, 1):
            cell = ws_line.cell(row=3, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border
        
        for row_idx, line in enumerate(line_data["lines"], 4):
            ws_line.cell(row=row_idx, column=1, value=line["line_name"]).border = thin_border
            ws_line.cell(row=row_idx, column=2, value=line["total_duration"]).border = thin_border
            ws_line.cell(row=row_idx, column=3, value=line["count"]).border = thin_border
            ws_line.cell(row=row_idx, column=4, value=line["avg_duration"]).border = thin_border
            ws_line.cell(row=row_idx, column=5, value=line["planned_duration"]).border = thin_border
            ws_line.cell(row=row_idx, column=6, value=line["unplanned_duration"]).border = thin_border
            ws_line.cell(row=row_idx, column=7, value=line["unplanned_ratio"]).border = thin_border
        
        for col in ["A", "B", "C", "D", "E", "F", "G"]:
            ws_line.column_dimensions[col].width = 18
    
    if efficiency_data.get("by_person"):
        ws_maint = wb.create_sheet("维修效率")
        ws_maint["A1"] = "维修人员效率分析"
        ws_maint["A1"].font = Font(bold=True, size=14, color="1E3A5F")
        ws_maint.merge_cells("A1:F1")
        
        maint_headers = ["维修人员", "技能等级", "团队", "平均修复时间(分钟)", "完成工单数量", "平均人工成本"]
        for col, header in enumerate(maint_headers, 1):
            cell = ws_maint.cell(row=3, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border
        
        for row_idx, person in enumerate(efficiency_data["by_person"], 4):
            ws_maint.cell(row=row_idx, column=1, value=person["person_name"]).border = thin_border
            ws_maint.cell(row=row_idx, column=2, value=person["skill_level"]).border = thin_border
            ws_maint.cell(row=row_idx, column=3, value=person["team"]).border = thin_border
            ws_maint.cell(row=row_idx, column=4, value=person["mttr"]).border = thin_border
            ws_maint.cell(row=row_idx, column=5, value=person["completed_count"]).border = thin_border
            ws_maint.cell(row=row_idx, column=6, value=person["avg_cost"]).border = thin_border
        
        for col in ["A", "B", "C", "D", "E", "F"]:
            ws_maint.column_dimensions[col].width = 18
    
    if correlation_data.get("matrix"):
        ws_parts = wb.create_sheet("备件关联")
        ws_parts["A1"] = "故障-备件关联分析"
        ws_parts["A1"].font = Font(bold=True, size=14, color="1E3A5F")
        ws_parts.merge_cells("A1:F1")
        
        parts_headers = ["故障类型", "备件名称", "使用次数", "总成本(元)", "关联度评分"]
        for col, header in enumerate(parts_headers, 1):
            cell = ws_parts.cell(row=3, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border
        
        sorted_matrix = sorted(correlation_data["matrix"], key=lambda x: x["correlation_score"], reverse=True)
        for row_idx, item in enumerate(sorted_matrix, 4):
            ws_parts.cell(row=row_idx, column=1, value=item["fault_name"]).border = thin_border
            ws_parts.cell(row=row_idx, column=2, value=item["part_name"]).border = thin_border
            ws_parts.cell(row=row_idx, column=3, value=item["usage_count"]).border = thin_border
            ws_parts.cell(row=row_idx, column=4, value=item["total_cost"]).border = thin_border
            ws_parts.cell(row=row_idx, column=5, value=item["correlation_score"]).border = thin_border
        
        for col in ["A", "B", "C", "D", "E"]:
            ws_parts.column_dimensions[col].width = 20
    
    ws_caliber = wb.create_sheet("数据口径")
    ws_caliber["A1"] = "数据口径说明"
    ws_caliber["A1"].font = Font(bold=True, size=14, color="1E3A5F")
    ws_caliber.merge_cells("A1:B1")
    
    caliber_headers = ["指标名称", "口径说明"]
    for col, header in enumerate(caliber_headers, 1):
        cell = ws_caliber.cell(row=3, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border
    
    caliber_items = [
        ["总停机时长", get_caliber_note("total_duration")],
        ["突发停机时长", get_caliber_note("unplanned_duration")],
        ["计划检修时长", get_caliber_note("planned_duration")],
        ["MTTR", get_caliber_note("mttr")],
        ["MTBF", get_caliber_note("mtbf")],
        ["设备可用率", get_caliber_note("availability")],
        ["Pareto分析", get_caliber_note("pareto")],
        ["停机类型划分", get_caliber_note("breakdown_type")],
    ]
    
    for row_idx, (name, note) in enumerate(caliber_items, 4):
        ws_caliber.cell(row=row_idx, column=1, value=name).border = thin_border
        ws_caliber.cell(row=row_idx, column=2, value=note).border = thin_border
        ws_caliber.cell(row=row_idx, column=2).alignment = Alignment(wrap_text=True)
    
    ws_caliber.column_dimensions["A"].width = 20
    ws_caliber.column_dimensions["B"].width = 60
    
    wb.save(output)
    output.seek(0)
    
    return output.getvalue()
