import json
from datetime import datetime, date
from typing import List, Optional, Any
from io import BytesIO

from sqlalchemy.orm import Session, Query
from openpyxl import Workbook

from app import models
from app.redis_client import RedisQueue


class DataScopeService:
    PROD_MODE = models.RunMode.PRODUCTION

    @staticmethod
    def is_prod_data(record) -> bool:
        return getattr(record, "run_mode", None) == DataScopeService.PROD_MODE

    @staticmethod
    def filter_real_reports(query: Query) -> Query:
        try:
            desc = query.column_descriptions
            if desc and len(desc) > 0:
                entity = desc[0].get("entity")
                if entity and hasattr(entity, "run_mode"):
                    return query.filter(entity.run_mode == DataScopeService.PROD_MODE)
        except Exception:
            pass
        try:
            return query.filter(
                getattr(query._entity_from_pre_ent_zero().class_, "run_mode", None)
                == DataScopeService.PROD_MODE
            )
        except Exception:
            return query.filter(
                models.BaseMixin.run_mode.property.columns[0] == DataScopeService.PROD_MODE
            )


class NotificationService:
    @staticmethod
    def create_notification(
        db: Session,
        user_id: int,
        title: str,
        content: str,
        ntype: str = "system",
        related_id: Optional[int] = None,
        related_type: Optional[str] = None,
        channel: str = "inapp",
        creator_id: Optional[int] = None,
    ) -> models.Notification:
        notif = models.Notification(
            user_id=user_id,
            title=title,
            content=content,
            notification_type=ntype,
            related_id=related_id,
            related_type=related_type,
            channel=channel,
            created_by=creator_id,
            updated_by=creator_id,
        )
        db.add(notif)
        db.flush()
        try:
            payload = json.dumps(
                {
                    "id": notif.id,
                    "user_id": user_id,
                    "title": title,
                    "content": content,
                    "type": ntype,
                    "created_at": notif.created_at.isoformat() if notif.created_at else None,
                },
                ensure_ascii=False,
            )
            RedisQueue.publish_notification(user_id, payload)
        except Exception:
            pass
        return notif

    @staticmethod
    def mark_read(db: Session, notif_id: int, user_id: int) -> Optional[models.Notification]:
        notif = (
            db.query(models.Notification)
            .filter(models.Notification.id == notif_id, models.Notification.user_id == user_id)
            .first()
        )
        if notif:
            notif.is_read = True
            notif.read_at = datetime.utcnow()
            notif.updated_by = user_id
            db.flush()
        return notif

    @staticmethod
    def mark_all_read(db: Session, user_id: int) -> int:
        updated = (
            db.query(models.Notification)
            .filter(
                models.Notification.user_id == user_id,
                models.Notification.is_read == False,
            )
            .update(
                {
                    models.Notification.is_read: True,
                    models.Notification.read_at: datetime.utcnow(),
                    models.Notification.updated_by: user_id,
                },
                synchronize_session=False,
            )
        )
        db.flush()
        return updated

    @staticmethod
    def list_unread(db: Session, user_id: int) -> List[models.Notification]:
        query = (
            db.query(models.Notification)
            .filter(models.Notification.user_id == user_id, models.Notification.is_read == False)
            .order_by(models.Notification.created_at.desc())
        )
        return DataScopeService.filter_real_reports(query).all()


class MonitorAlertService:
    @staticmethod
    def check_thresholds(
        db: Session,
        plot_id: int,
        variety_id: int,
        metrics_dict: dict,
        harvest_id: Optional[int] = None,
        creator_id: Optional[int] = None,
    ) -> List[models.Alert]:
        alerts: List[models.Alert] = []
        if not metrics_dict:
            return alerts

        thresholds = (
            db.query(models.Threshold)
            .filter(
                models.Threshold.variety_id == variety_id,
                models.Threshold.is_active == True,
            )
            .all()
        )
        thresholds = [t for t in thresholds if DataScopeService.is_prod_data(t)]

        for threshold in thresholds:
            metric = threshold.metric
            if metric not in metrics_dict or metrics_dict[metric] is None:
                continue

            actual_value = float(metrics_dict[metric])
            is_exceeded = False
            message_parts: List[str] = []

            if threshold.min_value is not None and actual_value < threshold.min_value:
                is_exceeded = True
                message_parts.append(f"{metric}={actual_value:.2f} 低于阈值下限 {threshold.min_value:.2f}")

            if threshold.max_value is not None and actual_value > threshold.max_value:
                is_exceeded = True
                message_parts.append(f"{metric}={actual_value:.2f} 超过阈值上限 {threshold.max_value:.2f}")

            if is_exceeded:
                message = "; ".join(message_parts) if message_parts else f"{metric}={actual_value:.2f} 超出阈值范围"
                alert = models.Alert(
                    harvest_id=harvest_id,
                    plot_id=plot_id,
                    variety_id=variety_id,
                    level=threshold.alert_level,
                    status=models.AlertStatus.OPEN,
                    metric=metric,
                    actual_value=actual_value,
                    threshold_min=threshold.min_value,
                    threshold_max=threshold.max_value,
                    message=message,
                    triggered_at=datetime.utcnow(),
                    created_by=creator_id,
                    updated_by=creator_id,
                )
                db.add(alert)
                db.flush()
                alerts.append(alert)

                try:
                    payload = json.dumps(
                        {
                            "id": alert.id,
                            "plot_id": plot_id,
                            "variety_id": variety_id,
                            "harvest_id": harvest_id,
                            "level": alert.level.value if hasattr(alert.level, "value") else str(alert.level),
                            "metric": metric,
                            "actual_value": actual_value,
                            "threshold_min": threshold.min_value,
                            "threshold_max": threshold.max_value,
                            "message": message,
                            "triggered_at": alert.triggered_at.isoformat(),
                        },
                        ensure_ascii=False,
                    )
                    RedisQueue.publish_alert(payload)
                except Exception:
                    pass

                try:
                    NotificationService.create_notification(
                        db=db,
                        user_id=creator_id if creator_id else 1,
                        title=f"[{alert.level.value if hasattr(alert.level, 'value') else str(alert.level).upper()}] 阈值告警",
                        content=message,
                        ntype="alert",
                        related_id=alert.id,
                        related_type="alert",
                        creator_id=creator_id,
                    )
                except Exception:
                    pass

        return alerts

    @staticmethod
    def acknowledge_alert(
        db: Session,
        alert_id: int,
        user_id: int,
        notes: Optional[str] = None,
    ) -> Optional[models.Alert]:
        alert = db.query(models.Alert).filter(models.Alert.id == alert_id).first()
        if alert:
            alert.status = models.AlertStatus.ACKNOWLEDGED
            alert.acknowledged_by = user_id
            alert.acknowledged_at = datetime.utcnow()
            if notes:
                alert.resolution_notes = notes
            alert.updated_by = user_id
            db.flush()
        return alert

    @staticmethod
    def resolve_alert(
        db: Session,
        alert_id: int,
        user_id: int,
        notes: Optional[str] = None,
    ) -> Optional[models.Alert]:
        alert = db.query(models.Alert).filter(models.Alert.id == alert_id).first()
        if alert:
            alert.status = models.AlertStatus.RESOLVED
            alert.resolved_by = user_id
            alert.resolved_at = datetime.utcnow()
            if notes:
                alert.resolution_notes = notes
            alert.updated_by = user_id
            db.flush()
        return alert


class YieldPredictionService:
    MODEL_VERSION = "v1.0-rule"

    @staticmethod
    def predict_yield(
        db: Session,
        harvest_record: models.HarvestRecord,
        creator_id: Optional[int] = None,
    ) -> models.YieldPrediction:
        variety = db.query(models.Variety).filter(models.Variety.id == harvest_record.variety_id).first()
        plot = db.query(models.Plot).filter(models.Plot.id == harvest_record.plot_id).first()

        area_mu = plot.area_mu if plot else 1.0
        expected_yield_kg = variety.expected_yield_kg if variety and variety.expected_yield_kg else 0.0

        env_factor = 1.0
        if harvest_record.temperature_c is not None:
            temp = harvest_record.temperature_c
            if 20 <= temp <= 30:
                env_factor *= 1.05
            elif temp < 15 or temp > 35:
                env_factor *= 0.85

        if harvest_record.humidity_pct is not None:
            hum = harvest_record.humidity_pct
            if 50 <= hum <= 80:
                env_factor *= 1.02
            elif hum < 30 or hum > 95:
                env_factor *= 0.9

        predicted_yield_kg: float
        confidence_pct: float
        features_used_parts: List[str] = []

        if expected_yield_kg and expected_yield_kg > 0:
            predicted_yield_kg = expected_yield_kg * area_mu * env_factor
            confidence_pct = 75.0
            features_used_parts.append("variety_expected_yield")
        else:
            predicted_yield_kg = area_mu * 500.0 * env_factor
            confidence_pct = 55.0
            features_used_parts.append("default_yield_per_mu")

        features_used_parts.extend(["plot_area", "env_temperature", "env_humidity"])
        confidence_pct = min(98.0, max(30.0, confidence_pct))

        harvest_record.predicted_yield_kg = predicted_yield_kg
        if harvest_record.actual_yield_kg and harvest_record.actual_yield_kg > 0:
            harvest_record.yield_deviation_pct = (
                (harvest_record.actual_yield_kg - predicted_yield_kg) / predicted_yield_kg * 100
            )

        prediction = models.YieldPrediction(
            harvest_id=harvest_record.id,
            plot_id=harvest_record.plot_id,
            variety_id=harvest_record.variety_id,
            prediction_date=harvest_record.harvest_date if isinstance(harvest_record.harvest_date, date) else date.today(),
            predicted_yield_kg=round(predicted_yield_kg, 2),
            confidence_pct=round(confidence_pct, 2),
            model_version=YieldPredictionService.MODEL_VERSION,
            features_used=",".join(features_used_parts),
            is_reminder_sent=False,
            created_by=creator_id,
            updated_by=creator_id,
        )
        db.add(prediction)
        db.flush()
        return prediction

    @staticmethod
    def send_reminder(
        db: Session,
        prediction: models.YieldPrediction,
    ) -> Optional[models.Notification]:
        if prediction.is_reminder_sent:
            return None

        harvest = db.query(models.HarvestRecord).filter(models.HarvestRecord.id == prediction.harvest_id).first()
        user_id = harvest.created_by if harvest and harvest.created_by else 1

        title = f"产量预测提醒 - 采收 #{prediction.harvest_id}"
        content = (
            f"预测产量: {prediction.predicted_yield_kg:.2f} kg, "
            f"置信度: {prediction.confidence_pct:.1f}%, "
            f"模型版本: {prediction.model_version}"
        )

        notif = NotificationService.create_notification(
            db=db,
            user_id=user_id,
            title=title,
            content=content,
            ntype="prediction",
            related_id=prediction.id,
            related_type="yield_prediction",
            creator_id=user_id,
        )

        prediction.is_reminder_sent = True
        prediction.reminder_sent_at = datetime.utcnow()
        prediction.updated_by = user_id
        db.flush()
        return notif


class HarvestChainService:
    @staticmethod
    def _parse_date(value: Any) -> date:
        if isinstance(value, date):
            return value
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, str):
            for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y%m%d"):
                try:
                    return datetime.strptime(value, fmt).date()
                except (ValueError, TypeError):
                    continue
        return date.today()

    @staticmethod
    def create_harvest(
        db: Session,
        data: dict,
        creator_id: int,
    ) -> models.HarvestRecord:
        run_auto_chain = bool(data.pop("run_auto_chain", True))

        harvest_date = HarvestChainService._parse_date(data.get("harvest_date"))
        status_value = data.get("status", "planned")
        if isinstance(status_value, str):
            try:
                status = models.HarvestStatus(status_value)
            except ValueError:
                status = models.HarvestStatus.PLANNED
        else:
            status = status_value

        harvest = models.HarvestRecord(
            code=data["code"],
            plot_id=data["plot_id"],
            variety_id=data["variety_id"],
            harvest_date=harvest_date,
            status=status,
            weather=data.get("weather"),
            temperature_c=data.get("temperature_c"),
            humidity_pct=data.get("humidity_pct"),
            workers_count=data.get("workers_count"),
            actual_yield_kg=data.get("actual_yield_kg"),
            predicted_yield_kg=data.get("predicted_yield_kg"),
            quality_score=data.get("quality_score"),
            notes=data.get("notes"),
            created_by=creator_id,
            updated_by=creator_id,
        )
        db.add(harvest)
        db.flush()

        if run_auto_chain:
            metrics_dict = {}
            if harvest.temperature_c is not None:
                metrics_dict["temperature_c"] = harvest.temperature_c
            if harvest.humidity_pct is not None:
                metrics_dict["humidity_pct"] = harvest.humidity_pct

            MonitorAlertService.check_thresholds(
                db=db,
                plot_id=harvest.plot_id,
                variety_id=harvest.variety_id,
                metrics_dict=metrics_dict,
                harvest_id=harvest.id,
                creator_id=creator_id,
            )

            prediction = YieldPredictionService.predict_yield(
                db=db,
                harvest_record=harvest,
                creator_id=creator_id,
            )
            YieldPredictionService.send_reminder(db=db, prediction=prediction)

        db.flush()
        return harvest

    @staticmethod
    def create_batch(
        db: Session,
        data: dict,
        creator_id: int,
    ) -> models.HarvestBatch:
        batch_date = HarvestChainService._parse_date(data.get("batch_date"))
        status_value = data.get("status", "harvested")
        if isinstance(status_value, str):
            try:
                status = models.BatchStatus(status_value)
            except ValueError:
                status = models.BatchStatus.HARVESTED
        else:
            status = status_value

        handler_id = data.get("handler_id", creator_id)

        batch = models.HarvestBatch(
            code=data["code"],
            harvest_id=data["harvest_id"],
            variety_id=data["variety_id"],
            batch_date=batch_date,
            status=status,
            weight_kg=data["weight_kg"],
            container_count=data.get("container_count"),
            storage_location=data.get("storage_location"),
            quality_grade=data.get("quality_grade"),
            handler_id=handler_id,
            handled_at=datetime.utcnow(),
            notes=data.get("notes"),
            created_by=creator_id,
            updated_by=creator_id,
        )
        db.add(batch)
        db.flush()
        return batch


class SubsidyBatchChainService:
    @staticmethod
    def create_subsidy_voucher(
        db: Session,
        data: dict,
        creator_id: int,
    ) -> models.SubsidyVoucher:
        voucher = models.SubsidyVoucher(
            code=data["code"],
            harvest_id=data["harvest_id"],
            batch_id=data.get("batch_id"),
            variety_id=data["variety_id"],
            subsidy_type=data["subsidy_type"],
            amount=float(data["amount"]),
            status=models.SubsidyStatus.PENDING,
            applicant_id=data.get("applicant_id", creator_id),
            applied_at=datetime.utcnow(),
            documents_ref=data.get("documents_ref"),
            review_notes=data.get("review_notes"),
            created_by=creator_id,
            updated_by=creator_id,
        )
        db.add(voucher)
        db.flush()

        try:
            applicant_id = voucher.applicant_id or creator_id
            NotificationService.create_notification(
                db=db,
                user_id=applicant_id,
                title=f"补贴凭证已提交 #{voucher.code}",
                content=f"已提交 {voucher.subsidy_type} 补贴申请，金额 ¥{voucher.amount:.2f}，等待审核。处理人：用户#{creator_id}，提交时间：{voucher.applied_at.strftime('%Y-%m-%d %H:%M')}",
                ntype="subsidy",
                related_id=voucher.id,
                related_type="subsidy_voucher",
                creator_id=creator_id,
            )
        except Exception:
            pass

        return voucher

    @staticmethod
    def review_subsidy(
        db: Session,
        voucher_id: int,
        reviewer_id: int,
        status: str,
        notes: Optional[str] = None,
    ) -> Optional[models.SubsidyVoucher]:
        voucher = db.query(models.SubsidyVoucher).filter(models.SubsidyVoucher.id == voucher_id).first()
        if not voucher:
            return None

        if isinstance(status, str):
            try:
                status_enum = models.SubsidyStatus(status)
            except ValueError:
                status_enum = models.SubsidyStatus.PENDING
        else:
            status_enum = status

        voucher.status = status_enum
        voucher.reviewer_id = reviewer_id
        voucher.reviewed_at = datetime.utcnow()
        if notes:
            voucher.review_notes = notes

        if status_enum == models.SubsidyStatus.PAID:
            voucher.paid_at = datetime.utcnow()

        voucher.updated_by = reviewer_id
        db.flush()

        try:
            status_text_map = {
                models.SubsidyStatus.PENDING: "待审核",
                models.SubsidyStatus.APPROVED: "已通过",
                models.SubsidyStatus.REJECTED: "已拒绝",
                models.SubsidyStatus.PAID: "已支付",
            }
            status_text = status_text_map.get(status_enum, str(status_enum))
            applicant_id = voucher.applicant_id or reviewer_id

            NotificationService.create_notification(
                db=db,
                user_id=applicant_id,
                title=f"补贴凭证审核状态 #{voucher.code}",
                content=f"补贴 {voucher.subsidy_type}（¥{voucher.amount:.2f}）状态更新为【{status_text}】。审核人：用户#{reviewer_id}，审核时间：{voucher.reviewed_at.strftime('%Y-%m-%d %H:%M')}。{notes or ''}",
                ntype="subsidy",
                related_id=voucher.id,
                related_type="subsidy_voucher",
                creator_id=reviewer_id,
            )
        except Exception:
            pass

        return voucher


class SortingService:
    @staticmethod
    def create_sorting_diff(
        db: Session,
        data: dict,
        creator_id: int,
    ) -> models.SortingDifference:
        reported_weight_kg = float(data["reported_weight_kg"])
        actual_weight_kg = float(data["actual_weight_kg"])
        diff_weight_kg = actual_weight_kg - reported_weight_kg
        diff_pct = (diff_weight_kg / reported_weight_kg * 100) if reported_weight_kg != 0 else 0.0

        result_value = data["result"]
        if isinstance(result_value, str):
            try:
                result = models.SortingResult(result_value)
            except ValueError:
                result = models.SortingResult.ACCEPTED
        else:
            result = result_value

        social_impact_value = data.get("social_impact", "none")
        if isinstance(social_impact_value, str):
            try:
                social_impact = models.SocialImpact(social_impact_value)
            except ValueError:
                social_impact = models.SocialImpact.NONE
        else:
            social_impact = social_impact_value

        handled_by = data.get("handled_by", creator_id)

        diff = models.SortingDifference(
            batch_id=data["batch_id"],
            harvest_id=data["harvest_id"],
            variety_id=data["variety_id"],
            reported_weight_kg=reported_weight_kg,
            actual_weight_kg=actual_weight_kg,
            diff_weight_kg=round(diff_weight_kg, 4),
            diff_pct=round(diff_pct, 4),
            reported_grade=data.get("reported_grade"),
            actual_grade=data.get("actual_grade"),
            remarks=data["remarks"],
            result=result,
            social_impact=social_impact,
            impact_description=data.get("impact_description"),
            handled_by=handled_by,
            handled_at=datetime.utcnow(),
            resolution=data.get("resolution"),
            created_by=creator_id,
            updated_by=creator_id,
        )
        db.add(diff)
        db.flush()

        try:
            result_text_map = {
                models.SortingResult.ACCEPTED: "正常接收",
                models.SortingResult.REJECTED: "拒收",
                models.SortingResult.REWORKED: "返工",
                models.SortingResult.DISCOUNTED: "折价处理",
            }
            result_text = result_text_map.get(result, str(result))
            social_impact_text_map = {
                models.SocialImpact.NONE: "无",
                models.SocialImpact.REPUTATION_RISK: "声誉风险",
                models.SocialImpact.CONSUMER_COMPLAINT: "消费者投诉",
                models.SocialImpact.REGULATORY_ATTENTION: "监管关注",
                models.SocialImpact.COMMUNITY_POSITIVE: "社区正面评价",
            }
            social_impact_text = social_impact_text_map.get(social_impact, str(social_impact))

            notify_user_ids = set()
            notify_user_ids.add(creator_id)
            notify_user_ids.add(handled_by)

            harvest = db.query(models.HarvestRecord).filter(models.HarvestRecord.id == diff.harvest_id).first()
            if harvest and harvest.created_by:
                notify_user_ids.add(harvest.created_by)

            for uid in notify_user_ids:
                if not uid:
                    continue
                NotificationService.create_notification(
                    db=db,
                    user_id=uid,
                    title=f"分选差异记录 #{diff.id}",
                    content=(
                        f"批次#{diff.batch_id} 分选差异："
                        f"上报重量 {reported_weight_kg:.2f}kg → 实际重量 {actual_weight_kg:.2f}kg "
                        f"（差异 {diff_weight_kg:+.2f}kg / {diff_pct:+.2f}%）。"
                        f"处理结果：【{result_text}】，"
                        f"社会影响：【{social_impact_text}】。"
                        f"处理人：用户#{handled_by}，"
                        f"处理时间：{diff.handled_at.strftime('%Y-%m-%d %H:%M')}。"
                        f"备注：{diff.remarks}"
                    ),
                    ntype="sorting",
                    related_id=diff.id,
                    related_type="sorting_difference",
                    creator_id=creator_id,
                )
        except Exception:
            pass

        return diff


class DownloadLogService:
    @staticmethod
    def log_download(
        db: Session,
        user_id: int,
        record_type: str,
        file_name: str,
        file_size: Optional[int] = None,
        filters: Optional[dict] = None,
        include_demo: bool = False,
    ) -> models.DownloadRecord:
        if isinstance(record_type, str):
            try:
                rt_enum = models.DownloadRecordType(record_type)
            except ValueError:
                rt_enum = models.DownloadRecordType.FULL_REPORT
        else:
            rt_enum = record_type

        filters_str = None
        if filters:
            try:
                filters_str = json.dumps(filters, ensure_ascii=False)
            except Exception:
                filters_str = str(filters)

        record = models.DownloadRecord(
            user_id=user_id,
            record_type=rt_enum,
            file_name=file_name,
            file_size_bytes=file_size,
            filters_used=filters_str,
            downloaded_at=datetime.utcnow(),
            include_demo=include_demo,
            created_by=user_id,
            updated_by=user_id,
        )
        db.add(record)
        db.flush()
        return record


class ReportExportService:
    @staticmethod
    def _build_query(db: Session, model, include_demo: bool = False) -> Query:
        query = db.query(model)
        if not include_demo:
            query = DataScopeService.filter_real_reports(query)
        return query

    @staticmethod
    def _workbook_to_bytes(wb: Workbook) -> bytes:
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.read()

    @staticmethod
    def _safe_date(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, (datetime, date)):
            return value.isoformat()
        return str(value)

    @staticmethod
    def export_harvests_excel(
        db: Session,
        start_date: Optional[date] = None,
        end_date: Optional[date] = None,
        include_demo: bool = False,
        user_id: Optional[int] = None,
    ) -> bytes:
        file_name = f"harvests_{start_date or 'all'}_{end_date or 'all'}.xlsx"
        DownloadLogService.log_download(
            db=db,
            user_id=user_id or 1,
            record_type=models.DownloadRecordType.HARVEST_REPORT,
            file_name=file_name,
            filters={"start_date": str(start_date), "end_date": str(end_date)},
            include_demo=include_demo,
        )

        query = ReportExportService._build_query(db, models.HarvestRecord, include_demo)
        if start_date:
            query = query.filter(models.HarvestRecord.harvest_date >= start_date)
        if end_date:
            query = query.filter(models.HarvestRecord.harvest_date <= end_date)
        records = query.order_by(models.HarvestRecord.harvest_date.desc()).all()

        wb = Workbook()
        ws = wb.active
        ws.title = "采收记录"
        ws.append(
            [
                "ID",
                "编码",
                "地块ID",
                "品种ID",
                "采收日期",
                "状态",
                "天气",
                "温度(℃)",
                "湿度(%)",
                "工人数",
                "实际产量(kg)",
                "预测产量(kg)",
                "偏差(%)",
                "质量评分",
                "备注",
                "创建人",
                "创建时间",
            ]
        )
        for r in records:
            ws.append(
                [
                    r.id,
                    r.code,
                    r.plot_id,
                    r.variety_id,
                    ReportExportService._safe_date(r.harvest_date),
                    r.status.value if hasattr(r.status, "value") else str(r.status),
                    r.weather or "",
                    r.temperature_c or "",
                    r.humidity_pct or "",
                    r.workers_count or "",
                    r.actual_yield_kg or "",
                    r.predicted_yield_kg or "",
                    r.yield_deviation_pct or "",
                    r.quality_score or "",
                    r.notes or "",
                    r.created_by or "",
                    ReportExportService._safe_date(r.created_at),
                ]
            )

        return ReportExportService._workbook_to_bytes(wb)

    @staticmethod
    def export_batches_excel(
        db: Session,
        start_date: Optional[date] = None,
        end_date: Optional[date] = None,
        include_demo: bool = False,
        user_id: Optional[int] = None,
    ) -> bytes:
        file_name = f"batches_{start_date or 'all'}_{end_date or 'all'}.xlsx"
        DownloadLogService.log_download(
            db=db,
            user_id=user_id or 1,
            record_type=models.DownloadRecordType.BATCH_REPORT,
            file_name=file_name,
            filters={"start_date": str(start_date), "end_date": str(end_date)},
            include_demo=include_demo,
        )

        query = ReportExportService._build_query(db, models.HarvestBatch, include_demo)
        if start_date:
            query = query.filter(models.HarvestBatch.batch_date >= start_date)
        if end_date:
            query = query.filter(models.HarvestBatch.batch_date <= end_date)
        records = query.order_by(models.HarvestBatch.batch_date.desc()).all()

        wb = Workbook()
        ws = wb.active
        ws.title = "采收批次"
        ws.append(
            [
                "ID",
                "编码",
                "采收ID",
                "品种ID",
                "批次日期",
                "状态",
                "重量(kg)",
                "容器数",
                "存储位置",
                "质量等级",
                "处理人ID",
                "处理时间",
                "备注",
                "创建人",
                "创建时间",
            ]
        )
        for r in records:
            ws.append(
                [
                    r.id,
                    r.code,
                    r.harvest_id,
                    r.variety_id,
                    ReportExportService._safe_date(r.batch_date),
                    r.status.value if hasattr(r.status, "value") else str(r.status),
                    r.weight_kg,
                    r.container_count or "",
                    r.storage_location or "",
                    r.quality_grade or "",
                    r.handler_id or "",
                    ReportExportService._safe_date(r.handled_at),
                    r.notes or "",
                    r.created_by or "",
                    ReportExportService._safe_date(r.created_at),
                ]
            )

        return ReportExportService._workbook_to_bytes(wb)

    @staticmethod
    def export_sorting_excel(
        db: Session,
        start_date: Optional[date] = None,
        end_date: Optional[date] = None,
        include_demo: bool = False,
        user_id: Optional[int] = None,
    ) -> bytes:
        file_name = f"sorting_{start_date or 'all'}_{end_date or 'all'}.xlsx"
        DownloadLogService.log_download(
            db=db,
            user_id=user_id or 1,
            record_type=models.DownloadRecordType.SORTING_REPORT,
            file_name=file_name,
            filters={"start_date": str(start_date), "end_date": str(end_date)},
            include_demo=include_demo,
        )

        query = ReportExportService._build_query(db, models.SortingDifference, include_demo)
        records = query.order_by(models.SortingDifference.created_at.desc()).all()

        wb = Workbook()
        ws = wb.active
        ws.title = "分选差异"
        ws.append(
            [
                "ID",
                "批次ID",
                "采收ID",
                "品种ID",
                "上报重量(kg)",
                "实际重量(kg)",
                "差异重量(kg)",
                "差异(%)",
                "上报等级",
                "实际等级",
                "备注",
                "结果",
                "社会影响",
                "影响描述",
                "处理人ID",
                "处理时间",
                "解决方案",
                "创建人",
                "创建时间",
            ]
        )
        for r in records:
            ws.append(
                [
                    r.id,
                    r.batch_id,
                    r.harvest_id,
                    r.variety_id,
                    r.reported_weight_kg,
                    r.actual_weight_kg,
                    r.diff_weight_kg,
                    r.diff_pct,
                    r.reported_grade or "",
                    r.actual_grade or "",
                    r.remarks,
                    r.result.value if hasattr(r.result, "value") else str(r.result),
                    r.social_impact.value if hasattr(r.social_impact, "value") else str(r.social_impact),
                    r.impact_description or "",
                    r.handled_by or "",
                    ReportExportService._safe_date(r.handled_at),
                    r.resolution or "",
                    r.created_by or "",
                    ReportExportService._safe_date(r.created_at),
                ]
            )

        return ReportExportService._workbook_to_bytes(wb)
