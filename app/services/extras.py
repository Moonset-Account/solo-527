from __future__ import annotations

import random
import re
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any, Optional
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import and_, desc, func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app import models
from app.core.config import settings
from app.core.security import hash_password
from app.db.redis_cache import RedisCache
from app.schemas import *  # noqa: F401,F403
from app.schemas import (
    AnomalyOut,
    AnomalySeverity,
    AnomalyStatus,
    AnomalyType,
    DashboardStats,
    ExportCreateIn,
    ExportStatus,
    ExportTaskOut,
    ExportType,
    ForecastItem,
    InvoiceCreateIn,
    InvoiceError,
    InvoiceOut,
    InvoiceStatus,
    InvoiceType,
    KPIData,
    MessageOut,
    PaginatedOut,
    PaginationIn,
    PrepaidAccountOut,
    PrepaidTxnOut,
    PrepaidTxnType,
    ReminderType,
    SessionUser,
    TodayTasks,
    TopAnomalyItem,
    TopClientItem,
    TrendPoint,
    AgingBucket,
)
from app.services import AnomalyService


def _to_decimal(v: Any) -> Decimal:
    return Decimal(str(v)) if v is not None else Decimal("0")


def _today() -> date:
    return datetime.now(timezone.utc).date()


def _paginate(total: int, items: list, p: PaginationIn) -> PaginatedOut:
    total_pages = (total + p.page_size - 1) // p.page_size if total > 0 else 1
    return PaginatedOut(items=items, total=total, page=p.page, page_size=p.page_size, total_pages=total_pages)


def _cn_tax_id_valid(tax_id: str) -> bool:
    if not tax_id or len(tax_id) not in (15, 17, 18, 20):
        return False
    return bool(re.match(r"^[0-9A-HJ-NPQRTUWXY]{15,20}$", tax_id.upper()))


class InvoiceService:
    def __init__(self, db: AsyncSession, operator: SessionUser, client_ip: str = ""):
        self.db = db
        self.operator = operator
        self.client_ip = client_ip
        self._anomaly_svc = AnomalyService(db, operator)

    def _validate(self, data: InvoiceCreateIn) -> list[InvoiceError]:
        errors: list[InvoiceError] = []
        tax_id = (data.tax_id or "").strip()
        if not _cn_tax_id_valid(tax_id):
            errors.append(InvoiceError(field="tax_id", error_code="INVALID_TAX_ID", message="纳税人识别号格式不正确", severity="ERROR"))
        title = (data.title or "").strip()
        if len(title) < 2:
            errors.append(InvoiceError(field="title", error_code="SHORT_TITLE", message="发票抬头过短", severity="WARNING"))
        if data.type in ("VAT_SPECIAL",):
            if not (data.address and data.phone):
                errors.append(InvoiceError(field="address/phone", error_code="MISSING_ADDR", message="专票需填写地址电话", severity="ERROR"))
            if not (data.bank_name and data.bank_account):
                errors.append(InvoiceError(field="bank", error_code="MISSING_BANK", message="专票需填写开户行及账号", severity="ERROR"))
        return errors

    async def list(self, p: PaginationIn, status: Optional[InvoiceStatus] = None) -> PaginatedOut[InvoiceOut]:
        stmt = select(models.Invoice).options(selectinload(models.Invoice.bill), selectinload(models.Invoice.client))
        if status:
            stmt = stmt.where(models.Invoice.status == status)
        stmt = stmt.order_by(desc(models.Invoice.applied_at))
        total = (await self.db.execute(select(func.count()).select_from(stmt.subquery()))).scalar_one() or 0
        stmt = stmt.offset(p.offset).limit(p.limit)
        rows = (await self.db.execute(stmt)).scalars().all()
        outs: list[InvoiceOut] = []
        for r in rows:
            o = InvoiceOut.model_validate(r)
            v = r.validation_errors
            if isinstance(v, dict) and "errors" in v:
                o.validation_errors = [InvoiceError.model_validate(e) for e in v["errors"]]
            outs.append(o)
        return _paginate(total, outs, p)

    async def apply(self, data: InvoiceCreateIn, client_id_override: Optional[UUID] = None) -> InvoiceOut:
        errors = self._validate(data)
        actual_client_id = client_id_override or self.operator.client_id
        inv = models.Invoice(
            bill_id=data.bill_id,
            client_id=actual_client_id,
            title=data.title,
            tax_id=data.tax_id,
            address=data.address,
            phone=data.phone,
            bank_name=data.bank_name,
            bank_account=data.bank_account,
            amount=data.amount,
            type=data.type,
            status="ERROR" if any(e.severity == "ERROR" for e in errors) else "PENDING",
            validation_errors={"errors": [e.model_dump() for e in errors]},
            applied_by=self.operator.id,
        )
        self.db.add(inv)
        await self.db.flush()
        bill_stmt = select(models.Bill).where(models.Bill.id == data.bill_id)
        bill = (await self.db.execute(bill_stmt)).scalar_one_or_none()
        if bill:
            bill.invoice_requested = True
        if errors:
            err_summary = "; ".join(f"{e.field}:{e.message}" for e in errors[:3])
            await self._anomaly_svc.create(
                type="INVOICE_ERROR",
                severity="HIGH" if any(e.severity == "ERROR" for e in errors) else "MEDIUM",
                title=f"发票信息错误：{data.title[:20]}",
                description=err_summary,
                related_entity_type="INVOICE",
                related_entity_id=inv.id,
            )
        await self.db.commit()
        await self.db.refresh(inv)
        out = InvoiceOut.model_validate(inv)
        v = inv.validation_errors
        if isinstance(v, dict) and "errors" in v:
            out.validation_errors = [InvoiceError.model_validate(e) for e in v["errors"]]
        return out

    async def update_status(self, inv_id: UUID, status: InvoiceStatus, invoice_no: Optional[str] = None, tracking_no: Optional[str] = None) -> Optional[InvoiceOut]:
        stmt = select(models.Invoice).where(models.Invoice.id == inv_id)
        inv = (await self.db.execute(stmt)).scalar_one_or_none()
        if not inv:
            return None
        inv.status = status
        if invoice_no:
            inv.invoice_no = invoice_no
        if tracking_no:
            inv.tracking_no = tracking_no
        now = datetime.now(timezone.utc)
        if status == "APPROVED":
            inv.approved_by = self.operator.id; inv.approved_at = now
        if status == "ISSUED":
            inv.issued_at = now
        if status == "MAILED":
            inv.mailed_at = now
        await self.db.commit()
        await self.db.refresh(inv)
        return InvoiceOut.model_validate(inv)


class PrepaidService:
    def __init__(self, db: AsyncSession, operator: SessionUser, client_ip: str = ""):
        self.db = db
        self.operator = operator
        self.client_ip = client_ip

    async def list_accounts(self, p: PaginationIn) -> PaginatedOut[PrepaidAccountOut]:
        stmt = select(models.PrepaidAccount).options(selectinload(models.PrepaidAccount.client)).order_by(desc(models.PrepaidAccount.updated_at))
        total = (await self.db.execute(select(func.count()).select_from(stmt.subquery()))).scalar_one() or 0
        stmt = stmt.offset(p.offset).limit(p.limit)
        rows = (await self.db.execute(stmt)).scalars().all()
        outs: list[PrepaidAccountOut] = []
        for a in rows:
            o = PrepaidAccountOut.model_validate(a)
            o.client_name = a.client.name if a.client else ""
            outs.append(o)
        return _paginate(total, outs, p)

    async def get_client_account(self, client_id: UUID) -> Optional[PrepaidAccountOut]:
        stmt = select(models.PrepaidAccount).options(selectinload(models.PrepaidAccount.client)).where(models.PrepaidAccount.client_id == client_id)
        a = (await self.db.execute(stmt)).scalar_one_or_none()
        if not a:
            return None
        out = PrepaidAccountOut.model_validate(a)
        out.client_name = a.client.name if a.client else ""
        return out

    async def list_txns(self, account_id: UUID, p: PaginationIn) -> PaginatedOut[PrepaidTxnOut]:
        stmt = select(models.PrepaidTransaction).where(models.PrepaidTransaction.account_id == account_id).order_by(desc(models.PrepaidTransaction.created_at))
        total = (await self.db.execute(select(func.count()).select_from(stmt.subquery()))).scalar_one() or 0
        stmt = stmt.offset(p.offset).limit(p.limit)
        rows = (await self.db.execute(stmt)).scalars().all()
        return _paginate(total, [PrepaidTxnOut.model_validate(r) for r in rows], p)

    async def recharge(self, client_id: UUID, amount: Decimal, remark: str = "") -> PrepaidAccountOut:
        stmt = select(models.PrepaidAccount).where(models.PrepaidAccount.client_id == client_id)
        account = (await self.db.execute(stmt)).scalar_one_or_none()
        client_stmt = select(models.Client).where(models.Client.id == client_id)
        client = (await self.db.execute(client_stmt)).scalar_one_or_none()
        if not account or not client:
            raise ValueError("账户不存在")
        before = account.current_balance
        account.current_balance = before + amount
        account.total_recharged = account.total_recharged + amount
        client.prepaid_balance = account.current_balance
        txn = models.PrepaidTransaction(
            account_id=account.id, type="RECHARGE", amount=amount,
            balance_before=before, balance_after=account.current_balance,
            operator_id=self.operator.id, remark=remark or "在线充值",
        )
        self.db.add(txn)
        await self.db.commit()
        await self.db.refresh(account)
        return await self.get_client_account(client_id) or PrepaidAccountOut.model_validate(account)

    async def deduct_for_bill(self, client_id: UUID, bill_id: UUID, amount: Decimal) -> tuple[Optional[PrepaidAccountOut], Optional[PrepaidTxnOut]]:
        stmt = select(models.PrepaidAccount).where(models.PrepaidAccount.client_id == client_id)
        account = (await self.db.execute(stmt)).scalar_one_or_none()
        bill_stmt = select(models.Bill).where(models.Bill.id == bill_id)
        bill = (await self.db.execute(bill_stmt)).scalar_one_or_none()
        client_stmt = select(models.Client).where(models.Client.id == client_id)
        client = (await self.db.execute(client_stmt)).scalar_one_or_none()
        if not account or not bill or not client:
            return None, None
        if account.current_balance < amount - Decimal("0.01"):
            return None, None
        before = account.current_balance
        deduct_amount = min(amount, before)
        account.current_balance = before - deduct_amount
        account.total_deducted = account.total_deducted + deduct_amount
        client.prepaid_balance = account.current_balance
        txn = models.PrepaidTransaction(
            account_id=account.id, type="DEDUCT", amount=deduct_amount,
            balance_before=before, balance_after=account.current_balance,
            related_bill_id=bill.id, operator_id=self.operator.id, remark="预存抵扣账单",
        )
        self.db.add(txn)
        pay = models.Payment(
            transaction_id=None, bill_id=bill.id, amount=deduct_amount,
            match_type="PREPAID", matched_by=self.operator.id, remark="预存抵扣",
        )
        self.db.add(pay)
        bill.paid_amount = bill.paid_amount + deduct_amount
        if bill.total_amount - bill.paid_amount <= Decimal("0.01"):
            bill.status = "PAID"
        elif bill.paid_amount > 0 and bill.status == "ISSUED":
            bill.status = "PARTIAL"
        await self.db.commit()
        await self.db.refresh(account)
        await self.db.refresh(txn)
        return PrepaidAccountOut.model_validate(account), PrepaidTxnOut.model_validate(txn)


class DashboardService:
    def __init__(self, db: AsyncSession, cache: RedisCache):
        self.db = db
        self.cache = cache

    async def get_stats(self) -> DashboardStats:
        cache_key = f"cache:dashboard:stats:{_today().isoformat()}"
        cached = await self.cache.get_json(cache_key)
        if cached:
            return DashboardStats.model_validate(cached)
        today = _today()
        month_start = today.replace(day=1)
        next_month = (month_start + timedelta(days=32)).replace(day=1)
        prev_month_start = (month_start - timedelta(days=1)).replace(day=1)

        # ===== 今日任务 =====
        t_stmt = (select(func.count(models.Bill.id),
                        func.coalesce(func.sum(models.Bill.total_amount - models.Bill.paid_amount), 0))
            .where(and_(models.Bill.due_date == today,
                        models.Bill.status.in_(["ISSUED", "PARTIAL", "OVERDUE"]))))
        t_count, t_amt = (await self.db.execute(t_stmt)).one()
        u_stmt = (select(func.count(models.BankTransaction.id),
                         func.coalesce(func.sum(models.BankTransaction.amount), 0))
            .where(models.BankTransaction.match_status == "UNMATCHED"))
        u_count, u_amt = (await self.db.execute(u_stmt)).one()
        i_stmt = (select(func.count(models.Invoice.id),
                         func.coalesce(func.sum(models.Invoice.amount), 0))
            .where(models.Invoice.status.in_(["PENDING", "APPROVED"])))
        i_count, i_amt = (await self.db.execute(i_stmt)).one()
        open_anom = (await self.db.execute(
            select(func.count(models.Anomaly.id)).where(models.Anomaly.status == "OPEN")
        )).scalar_one() or 0
        tasks = TodayTasks(
            today_bills_count=int(t_count or 0),
            today_due_amount=_to_decimal(t_amt),
            unmatched_txns_count=int(u_count or 0),
            unmatched_txns_amount=_to_decimal(u_amt),
            pending_invoices_count=int(i_count or 0),
            pending_invoices_amount=_to_decimal(i_amt),
            open_anomalies_count=int(open_anom),
        )

        # ===== KPI =====
        outstanding = (await self.db.execute(
            select(func.coalesce(func.sum(models.Bill.total_amount - models.Bill.paid_amount), 0))
            .where(models.Bill.status.notin_(["PAID", "VOID"]))
        )).scalar_one() or 0
        prev_outstanding = (await self.db.execute(
            select(func.coalesce(func.sum(models.Bill.total_amount - models.Bill.paid_amount), 0))
            .where(and_(models.Bill.status.notin_(["PAID", "VOID"]), models.Bill.issue_date < month_start))
        )).scalar_one() or 0
        os_change = (
            round(float((_to_decimal(outstanding) - _to_decimal(prev_outstanding)) /
                        max(_to_decimal(prev_outstanding), Decimal("1")) * 100), 2)
        )
        collected = (await self.db.execute(
            select(func.coalesce(func.sum(models.Payment.amount), 0))
            .where(and_(models.Payment.matched_at >= month_start, models.Payment.matched_at < next_month))
        )).scalar_one() or 0
        prev_collected = (await self.db.execute(
            select(func.coalesce(func.sum(models.Payment.amount), 0))
            .where(and_(models.Payment.matched_at >= prev_month_start, models.Payment.matched_at < month_start))
        )).scalar_one() or 0
        collected_change = (
            round(float((_to_decimal(collected) - _to_decimal(prev_collected)) /
                        max(_to_decimal(prev_collected), Decimal("1")) * 100), 2)
        )
        total_billed = (await self.db.execute(
            select(func.coalesce(func.sum(models.Bill.total_amount), 0))
            .where(and_(models.Bill.issue_date >= month_start, models.Bill.issue_date < next_month))
        )).scalar_one() or 0
        collection_rate = (
            round(float(_to_decimal(collected) / max(_to_decimal(total_billed), Decimal("1")) * 100), 2)
        )
        overdue_amt = (await self.db.execute(
            select(func.coalesce(func.sum(models.Bill.total_amount - models.Bill.paid_amount), 0))
            .where(and_(models.Bill.days_overdue > 0, models.Bill.status != "PAID"))
        )).scalar_one() or 0
        overdue_rate = (
            round(float(_to_decimal(overdue_amt) / max(_to_decimal(outstanding), Decimal("1")) * 100), 2)
        )
        kpi = KPIData(
            total_outstanding=_to_decimal(outstanding),
            total_outstanding_change_pct=os_change,
            collected_this_month=_to_decimal(collected),
            collected_change_pct=collected_change,
            collection_rate=collection_rate,
            overdue_rate=overdue_rate,
            total_anomalies=int(open_anom),
        )

        # ===== 账龄分布 =====
        aging: dict[str, AgingBucket] = {}
        for label, lo, hi in [("0-30", 0, 30), ("31-60", 31, 60), ("61-90", 61, 90), ("90+", 91, 9999999)]:
            c, a = (await self.db.execute(select(
                func.count(models.Bill.id),
                func.coalesce(func.sum(models.Bill.total_amount - models.Bill.paid_amount), 0),
            ).where(and_(
                models.Bill.status != "PAID", models.Bill.days_overdue >= lo, models.Bill.days_overdue <= hi
            )))).one()
            aging[label] = AgingBucket(count=int(c or 0), amount=_to_decimal(a))

        # ===== 趋势 6 个月 =====
        trend: list[TrendPoint] = []
        for i in range(5, -1, -1):
            ref = today - timedelta(days=i * 30)
            m_start = ref.replace(day=1)
            m_end = (m_start + timedelta(days=32)).replace(day=1)
            m_bill = (await self.db.execute(select(func.coalesce(func.sum(models.Bill.total_amount), 0))
                .where(and_(models.Bill.issue_date >= m_start, models.Bill.issue_date < m_end)))).scalar_one() or 0
            m_coll = (await self.db.execute(select(func.coalesce(func.sum(models.Payment.amount), 0))
                .where(and_(models.Payment.matched_at >= m_start, models.Payment.matched_at < m_end)))).scalar_one() or 0
            rt = (
                round(float(_to_decimal(m_coll) / max(_to_decimal(m_bill), Decimal("1")) * 100), 2)
            )
            trend.append(TrendPoint(
                month=m_start.strftime("%y-%m"),
                billed=_to_decimal(m_bill),
                collected=_to_decimal(m_coll),
                rate=rt,
            ))

        # ===== 异常计数 + Top异常 =====
        anomaly_counts: dict[str, int] = {}
        for t, in (await self.db.execute(
            select(models.Anomaly.type).where(models.Anomaly.status == "OPEN")
        )).all():
            anomaly_counts[t] = anomaly_counts.get(t, 0) + 1
        top_anom_stmt = (select(models.Anomaly)
            .where(models.Anomaly.status == "OPEN")
            .order_by(desc(models.Anomaly.severity), desc(models.Anomaly.created_at))
            .limit(12))
        top_anom_rows = (await self.db.execute(top_anom_stmt)).scalars().all()
        top_anomalies: list[TopAnomalyItem] = []
        for a in top_anom_rows:
            # entity_no: 按类型从相关表取
            en = a.related_entity_id
            amt: Optional[Decimal] = None
            anom_desc = a.title
            if a.type in ("TXN_UNMATCHED", "AMOUNT_DIFF"):
                try:
                    if en:
                        r = (await self.db.execute(
                            select(models.BankTransaction.txn_no, models.BankTransaction.amount,
                                   models.BankTransaction.counterparty_name)
                            .where(models.BankTransaction.id == en))).one_or_none()
                        if r:
                            en_str, amt_r, cp = r
                            en = en_str
                            amt = _to_decimal(amt_r)
                            anom_desc = cp or a.title or a.description or ""
                except Exception:
                    pass
            elif a.type == "INVOICE_ERROR":
                try:
                    if en:
                        r = (await self.db.execute(
                            select(models.Invoice.invoice_no, models.Invoice.amount)
                            .where(models.Invoice.id == en))).one_or_none()
                        if r:
                            en_str, amt_r = r
                            en = en_str
                            amt = _to_decimal(amt_r)
                            anom_desc = (a.description or a.title or "")
                except Exception:
                    pass
            elif a.type == "OVERDUE" or a.type.startswith("OVERDUE"):
                try:
                    if en:
                        r = (await self.db.execute(
                            select(models.Bill.bill_no, models.Bill.total_amount - models.Bill.paid_amount,
                                   models.Client.name)
                            .join(models.Client, models.Client.id == models.Bill.client_id)
                            .where(models.Bill.id == en))).one_or_none()
                        if r:
                            en_str, amt_r, cname = r
                            en = en_str
                            amt = _to_decimal(amt_r)
                            anom_desc = cname
                except Exception:
                    pass
            sev = None
            try:
                sev_map = {"LOW": "低", "MEDIUM": "中", "HIGH": "高", "CRITICAL": "严重"}
                sev = sev_map.get(str(a.severity), str(a.severity))
            except Exception:
                pass
            top_anomalies.append(TopAnomalyItem(
                type=str(a.type), severity=sev, entity_no=(str(en) if en else None),
                amount=amt, description=(anom_desc or "")[:60],
                created_at=a.created_at.isoformat() if a.created_at else None,
            ))

        # ===== Top 5 客户回款 =====
        top_clients_stmt = (select(models.Client.name, func.coalesce(func.sum(models.Payment.amount), 0))
            .join(models.Bill, models.Bill.id == models.Payment.bill_id)
            .join(models.Client, models.Client.id == models.Bill.client_id)
            .where(and_(models.Payment.matched_at >= month_start, models.Payment.matched_at < next_month))
            .group_by(models.Client.name)
            .order_by(desc(func.sum(models.Payment.amount)))
            .limit(5))
        top_clients_rows = (await self.db.execute(top_clients_stmt)).all()
        top_clients: list[TopClientItem] = []
        for nm, am in top_clients_rows:
            top_clients.append(TopClientItem(client_name=nm or "未知", amount=_to_decimal(am)))

        # ===== 现金预测 30/60/90 =====
        async def _forecast(days: int) -> ForecastItem:
            horizon = today + timedelta(days=days)
            cnt_stmt = (select(func.count(models.Bill.id),
                               func.coalesce(func.sum(models.Bill.total_amount - models.Bill.paid_amount), 0))
                .where(and_(models.Bill.status != "PAID", models.Bill.due_date <= horizon)))
            bc, tot = (await self.db.execute(cnt_stmt)).one()
            od_stmt = (select(func.coalesce(func.sum(models.Bill.total_amount - models.Bill.paid_amount), 0))
                .where(and_(models.Bill.status != "PAID", models.Bill.days_overdue > 0, models.Bill.due_date <= horizon)))
            od = (await self.db.execute(od_stmt)).scalar_one() or 0
            po_stmt = (select(func.coalesce(func.sum(models.PrepaidAccount.current_balance), 0)))
            po = (await self.db.execute(po_stmt)).scalar_one() or 0
            offset = _to_decimal(po) * Decimal("0.6")
            return ForecastItem(
                total_expected=_to_decimal(tot),
                bills_count=int(bc or 0),
                overdue_expected=_to_decimal(od),
                prepaid_offset=offset,
            )
        f30 = await _forecast(30)
        f60 = await _forecast(60)
        f90 = await _forecast(90)

        stats = DashboardStats(
            _today=today.isoformat(),
            tasks=tasks, kpi=kpi, aging_buckets=aging,
            monthly_trend=trend, anomaly_counts=anomaly_counts,
            top_anomalies=top_anomalies, top_clients_collected=top_clients,
            forecast_30=f30, forecast_60=f60, forecast_90=f90,
        )
        await self.cache.set_json(cache_key, stats.model_dump(mode="json"), ttl=1800)
        return stats


class ExportService:
    def __init__(self, db: AsyncSession, cache: RedisCache, operator: SessionUser):
        self.db = db
        self.cache = cache
        self.operator = operator

    async def list_tasks(self, p: PaginationIn) -> PaginatedOut[ExportTaskOut]:
        stmt = select(models.ExportTask).where(models.ExportTask.created_by == self.operator.id).order_by(desc(models.ExportTask.created_at))
        total = (await self.db.execute(select(func.count()).select_from(stmt.subquery()))).scalar_one() or 0
        stmt = stmt.offset(p.offset).limit(p.limit)
        rows = (await self.db.execute(stmt)).scalars().all()
        return _paginate(total, [ExportTaskOut.model_validate(r) for r in rows], p)

    async def create_task(self, data: ExportCreateIn) -> ExportTaskOut:
        task = models.ExportTask(
            type=data.type, status="PENDING", parameters=data.parameters or {},
            created_by=self.operator.id, expires_at=datetime.now(timezone.utc) + timedelta(days=7),
        )
        self.db.add(task)
        await self.db.commit()
        await self.db.refresh(task)
        return ExportTaskOut.model_validate(task)

    async def _rows_for_cashflow(self) -> list[dict]:
        stmt = (select(models.Bill, models.Client.name, models.Payment)
            .join(models.Client, models.Bill.client_id == models.Client.id)
            .outerjoin(models.Payment, models.Payment.bill_id == models.Bill.id)
            .order_by(desc(models.Bill.due_date)).limit(5000))
        rows = (await self.db.execute(stmt)).all()
        out: list[dict] = []
        for bill, client_name, pay in rows:
            out.append({
                "账单编号": bill.bill_no, "客户名称": client_name,
                "账期": f"{bill.period_start} ~ {bill.period_end}",
                "出账日期": str(bill.issue_date), "到期日期": str(bill.due_date),
                "应收金额": float(bill.total_amount), "已收金额": float(bill.paid_amount),
                "待收金额": float(bill.total_amount - bill.paid_amount),
                "状态": bill.status, "逾期天数": bill.days_overdue,
                "最近收款日期": str(pay.matched_at.date()) if pay else "",
                "最近收款金额": float(pay.amount) if pay else 0,
            })
        return out

    async def _rows_for_invoice_errors(self) -> list[dict]:
        stmt = select(models.Invoice).order_by(desc(models.Invoice.applied_at)).limit(5000)
        rows = (await self.db.execute(stmt)).scalars().all()
        out: list[dict] = []
        for inv in rows:
            v = inv.validation_errors
            errs = v.get("errors", []) if isinstance(v, dict) else []
            for e in errs:
                out.append({
                    "发票ID": str(inv.id), "账单ID": str(inv.bill_id),
                    "发票抬头": inv.title, "税号": inv.tax_id,
                    "错误字段": e.get("field", ""), "错误编码": e.get("error_code", ""),
                    "错误信息": e.get("message", ""), "严重程度": e.get("severity", ""),
                    "申请时间": str(inv.applied_at), "当前状态": inv.status,
                })
        return out

    async def _rows_for_changelog(self, since_days: int = 30) -> list[dict]:
        since = datetime.now(timezone.utc) - timedelta(days=since_days)
        stmt = select(models.AuditLog).where(models.AuditLog.created_at >= since).order_by(desc(models.AuditLog.created_at)).limit(5000)
        rows = (await self.db.execute(stmt)).scalars().all()
        return [{
            "实体类型": r.entity_type, "实体ID": str(r.entity_id), "变更字段": r.field_name,
            "原值": r.old_value or "", "新值": r.new_value or "", "操作人": r.operator_name,
            "操作IP": r.operator_ip or "", "变更原因": r.change_reason or "", "操作时间": str(r.created_at),
        } for r in rows]

    async def execute_task(self, task_id: UUID) -> Optional[Path]:
        stmt = select(models.ExportTask).where(models.ExportTask.id == task_id)
        task = (await self.db.execute(stmt)).scalar_one_or_none()
        if not task:
            return None
        task.status = "PROCESSING"
        task.started_at = datetime.now(timezone.utc)
        await self.db.flush()
        rows: list[dict] = []
        if task.type == "CASHFLOW":
            rows = await self._rows_for_cashflow()
        elif task.type == "INVOICE_ERRORS":
            rows = await self._rows_for_invoice_errors()
        elif task.type == "CHANGE_LOG":
            days = 30
            if isinstance(task.parameters, dict):
                days = int(task.parameters.get("days", 30))
            rows = await self._rows_for_changelog(days)
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = {"CASHFLOW": "现金流报表", "INVOICE_ERRORS": "发票错误清单", "CHANGE_LOG": "最近变更记录"}.get(task.type, "导出数据")
            headers = list(rows[0].keys()) if rows else ["空"]
            ws.append(headers)
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill("solid", fgColor="2D7A5E")
            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=1, column=col_idx)
                c.font = header_font
                c.fill = header_fill
                c.alignment = Alignment(horizontal="center")
            for r in rows:
                ws.append([r.get(h, "") for h in headers])
            for col_idx in range(1, len(headers) + 1):
                ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = 18
            file_name = f"{task.type}_{task.id}.xlsx"
            file_path = settings.export_path / file_name
            wb.save(file_path)
            task.status = "COMPLETED"
            task.file_name = file_name
            task.file_size = file_path.stat().st_size
            task.row_count = len(rows)
            task.completed_at = datetime.now(timezone.utc)
            await self.db.commit()
            return file_path
        except Exception as e:
            task.status = "FAILED"
            task.error_message = str(e)
            await self.db.commit()
            return None

    async def _mark_failed(self, task_id: UUID, error: str) -> None:
        stmt = select(models.ExportTask).where(models.ExportTask.id == task_id)
        task = (await self.db.execute(stmt)).scalar_one_or_none()
        if task:
            task.status = "FAILED"
            task.error_message = error
            task.completed_at = datetime.now(timezone.utc)
            await self.db.flush()

    async def retry_task(self, task_id: UUID) -> Optional[ExportTaskOut]:
        stmt = select(models.ExportTask).where(and_(
            models.ExportTask.id == task_id, models.ExportTask.created_by == self.operator.id,
        ))
        task = (await self.db.execute(stmt)).scalar_one_or_none()
        if not task:
            return None
        task.status = "PENDING"
        task.started_at = None
        task.completed_at = None
        task.error_message = None
        task.file_name = None
        task.file_size = None
        task.row_count = None
        await self.db.commit()
        await self.db.refresh(task)
        return ExportTaskOut.model_validate(task)

    async def get_download_path(self, task_id: UUID) -> Optional[Path]:
        stmt = select(models.ExportTask).where(and_(
            models.ExportTask.id == task_id, models.ExportTask.created_by == self.operator.id,
            models.ExportTask.status == "COMPLETED",
        ))
        task = (await self.db.execute(stmt)).scalar_one_or_none()
        if not task or not task.file_name:
            return None
        p = settings.export_path / task.file_name
        return p if p.exists() else None


class MockDataService:
    def __init__(self, db: AsyncSession):
        self.db = db

    async def seed_all(self) -> MessageOut:
        cnt = (await self.db.execute(select(func.count()).select_from(models.User))).scalar_one() or 0
        if cnt > 2:
            return MessageOut(message="已有数据，跳过Mock生成", code=1)
        now = datetime.now(timezone.utc)
        today = now.date()
        clients_data = [
            ("北京禾田科技有限公司", "91110000MA001ABCD1", "张三", "13800138001", Decimal("120000")),
            ("上海青蓝信息技术有限公司", "91310000MA1FL2XY3Z", "李四", "13800138002", Decimal("85000")),
            ("深圳绿源环保科技股份有限公司", "91440300MA5EKLMN7P", "王五", "13800138003", Decimal("300000")),
            ("广州金穗农业发展有限公司", "91440101MA9UPQRS8Q", "赵六", "13800138004", Decimal("2000")),
            ("杭州云栖网络科技有限公司", "91330106MA2ABCD123", "钱七", "13800138005", Decimal("45000")),
            ("南京青松建筑工程有限公司", "91320100MA1NABCD12", "孙八", "13800138006", Decimal("15000")),
        ]
        clients: list[models.Client] = []
        for name, tax_id, contact, phone, balance in clients_data:
            c = models.Client(
                name=name, tax_id=tax_id, contact_name=contact, contact_phone=phone,
                address=f"{name[:2]}市示例区示例路{random.randint(1,999)}号",
                bank_name="中国工商银行示例支行", bank_account="622202000000" + str(random.randint(10000000, 99999999)),
                credit_limit=Decimal("100000"), prepaid_balance=balance,
            )
            self.db.add(c); clients.append(c)
        await self.db.flush()
        for idx, c in enumerate(clients):
            u = models.User(
                email=f"client{idx+1}@example.com", password_hash=hash_password("client123"),
                name=c.contact_name or f"客户{idx+1}", role="CLIENT", client_id=c.id,
            )
            self.db.add(u)
            acc = models.PrepaidAccount(
                client_id=c.id, current_balance=c.prepaid_balance,
                total_recharged=c.prepaid_balance + Decimal("50000"),
                total_deducted=Decimal("50000"),
            )
            self.db.add(acc)
        await self.db.flush()

        products = ["企业版订阅年费", "专业版订阅季费", "旗舰版订阅月费", "基础版订阅年费", "增值服务包"]
        for ci, c in enumerate(clients):
            for j in range(3):
                self.db.add(models.Order(
                    client_id=c.id, order_no=f"ORD{now.year}{ci+1:02d}{j+1:04d}",
                    product_name=products[(ci + j) % len(products)],
                    plan_type=["年费", "季费", "月费"][j % 3],
                    amount=Decimal(random.choice([12000, 3600, 980, 24000, 68000, 5600])),
                    start_date=today - timedelta(days=365 + j * 30),
                    end_date=today + timedelta(days=365 - j * 30),
                    auto_renew=True, status="ACTIVE",
                ))
        await self.db.flush()

        statuses: list = ["ISSUED", "ISSUED", "PARTIAL", "PAID", "OVERDUE", "OVERDUE", "PAID", "ISSUED"]
        for ci, c in enumerate(clients):
            for j in range(6):
                due_offset = random.choice([-45, -20, -5, 0, 10, 25, 40, 60])
                issue_offset = due_offset - 15
                period_start = today + timedelta(days=issue_offset - 30)
                period_end = today + timedelta(days=issue_offset)
                issue_date = today + timedelta(days=issue_offset)
                due_date = today + timedelta(days=due_offset)
                amt = Decimal(random.randint(2000, 80000))
                st: str = random.choice(statuses)
                paid = amt if st == "PAID" else amt * Decimal(random.uniform(0.1, 0.6)) if st == "PARTIAL" else Decimal("0")
                self.db.add(models.Bill(
                    bill_no=f"BL{period_start.strftime('%Y%m')}{ci+1:02d}{j+1:02d}",
                    client_id=c.id, order_id=None,
                    period_start=period_start, period_end=period_end,
                    issue_date=issue_date, due_date=due_date,
                    total_amount=amt, paid_amount=paid, status=st,
                    days_overdue=max(0, (today - due_date).days) if st != "PAID" else 0,
                    invoice_requested=random.choice([True, False, False, False]),
                    forecast_impact_score=Decimal(random.uniform(-0.5, 1)),
                ))
        await self.db.flush()

        for k in range(22):
            party = random.choice(clients)
            amt = Decimal(random.randint(3000, 90000))
            offset = random.randint(-30, 0)
            self.db.add(models.BankTransaction(
                txn_no=f"TXN{now.strftime('%Y%m%d')}{random.randint(1000, 9999)}{k}",
                txn_date=today + timedelta(days=offset), amount=amt, direction="IN",
                counterparty=party.name, counterparty_account=str(random.randint(600000, 999999)),
                summary=f"货款 订阅费 {random.choice(['', 'BL2024', '服务费', 'Q3'])}",
                balance_after=Decimal(random.randint(100000, 5000000)),
                match_status=random.choice(["UNMATCHED", "MATCHED", "PARTIAL", "UNMATCHED"]),
                matched_amount=Decimal("0"), confidence_score=None,
                import_batch=f"BATCH{now.strftime('%Y%m%d')}",
            ))
        await self.db.flush()

        inv_statuses: list = ["PENDING", "APPROVED", "ISSUING", "ISSUED", "ERROR", "PENDING"]
        for ci, c in enumerate(clients):
            bill_stmt = select(models.Bill).where(models.Bill.client_id == c.id).limit(2)
            bills_for_inv = (await self.db.execute(bill_stmt)).scalars().all()
            for b in bills_for_inv:
                st = random.choice(inv_statuses)
                errors: list[dict] = []
                if st == "ERROR":
                    errors = [{"field": "tax_id", "error_code": "INVALID_TAX_ID", "message": "纳税人识别号格式错误", "severity": "ERROR"}]
                inv = models.Invoice(
                    invoice_no=f"INV{now.strftime('%Y%m')}{random.randint(1000,9999)}" if st in ("ISSUED", "MAILED") else None,
                    bill_id=b.id, client_id=c.id, title=c.name, tax_id=c.tax_id or "",
                    address=c.address, phone=c.contact_phone,
                    bank_name=c.bank_name, bank_account=c.bank_account,
                    amount=b.total_amount, type=random.choice(["VAT_SPECIAL", "VAT_NORMAL", "ELECTRONIC"]),
                    status=st, validation_errors={"errors": errors},
                    applied_by=None, applied_at=now - timedelta(days=random.randint(0, 20)),
                )
                self.db.add(inv)
                if not b.invoice_requested:
                    b.invoice_requested = True
        await self.db.flush()

        anomaly_defs = [
            ("TXN_UNMATCHED", "流水无法匹配：对方户名模糊，需人工确认", "MEDIUM"),
            ("INVOICE_ERROR", "发票税号校验失败：长度不符合标准", "HIGH"),
            ("AMOUNT_DIFF", "账单与流水金额差异超过 5%", "MEDIUM"),
            ("OVERDUE_LIMIT", "账单逾期已超 30 天，建议升级处理", "CRITICAL"),
            ("FORECAST_WARN", "现金预测：下月回款置信度偏低", "LOW"),
            ("TXN_UNMATCHED", "大额流水（>5万）暂未匹配", "HIGH"),
        ]
        for at, desc, sev in anomaly_defs:
            self.db.add(models.Anomaly(
                type=at, severity=sev, title=desc[:30], description=desc,
                related_entity_type=None, related_entity_id=None, status="OPEN", assignee_id=None,
            ))

        reminder_tmpl = [
            ("BEFORE_DUE", 7, "EMAIL"),
            ("BEFORE_DUE", 3, "EMAIL"),
            ("BEFORE_DUE", 1, "EMAIL"),
            ("OVERDUE_DAILY", None, "IN_APP"),
        ]
        for type_name, days, ch in reminder_tmpl:
            bills = (await self.db.execute(select(models.Bill).where(models.Bill.status != "PAID").limit(5))).scalars().all()
            for b in bills[:2]:
                self.db.add(models.Reminder(
                    bill_id=b.id, type=type_name, days_before=days, channel=ch,
                    status=random.choice(["PENDING", "SENT", "SENT"]),
                    sent_at=now - timedelta(days=1) if random.random() > 0.5 else None,
                ))

        await self.db.commit()
        return MessageOut(message=f"Mock数据初始化完成：{len(clients)}个客户，约{6*len(clients)}条账单，22条流水")
