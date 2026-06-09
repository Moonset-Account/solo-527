import os
import re
import hashlib
import logging
from typing import List, Optional, Dict, Any, Tuple
from dataclasses import dataclass, field
from pathlib import Path
from datetime import datetime

from loguru import logger


@dataclass
class ParsedClause:
    index: int
    title: Optional[str] = None
    number: Optional[str] = None
    text: str = ""
    page_start: int = 1
    page_end: int = 1
    char_start: int = 0
    char_end: int = 0
    raw_category: Optional[str] = None
    metadata: Dict[str, Any] = field(default_factory=dict)


@dataclass
class ParsedDocument:
    file_name: str
    file_path: str
    file_size: int
    file_hash: str
    mime_type: str
    title: str = ""
    full_text: str = ""
    page_count: int = 0
    word_count: int = 0
    clauses: List[ParsedClause] = field(default_factory=list)
    metadata: Dict[str, Any] = field(default_factory=dict)
    parties: Dict[str, str] = field(default_factory=dict)
    key_dates: Dict[str, Optional[datetime]] = field(default_factory=dict)
    total_amount: Optional[float] = None
    currency: str = "CNY"
    error_message: Optional[str] = None


class DocumentParser:
    """
    合同文档解析器
    支持 PDF、DOCX、TXT 格式
    功能：文本提取、页码定位、条款切分、元数据抽取
    """

    SUPPORTED_MIME = {
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".txt": "text/plain",
    }

    CLAUSE_PATTERNS = [
        re.compile(r"^第[一二三四五六七八九十百千万\d]+[章节条条款项]", re.MULTILINE),
        re.compile(r"^(\d+[\.、\)]+)\s*", re.MULTILINE),
        re.compile(r"^([一二三四五六七八九十]+[\.、\)]+)\s*", re.MULTILINE),
        re.compile(r"^Article\s+\d+", re.MULTILINE | re.IGNORECASE),
        re.compile(r"^Section\s+\d+", re.MULTILINE | re.IGNORECASE),
        re.compile(r"^Clause\s+\d+", re.MULTILINE | re.IGNORECASE),
    ]

    TITLE_PATTERNS = [
        (re.compile(r"合同名称[:：]\s*(.+)"), "contract_name"),
        (re.compile(r"合同编号[:：]\s*(.+)"), "contract_no"),
        (re.compile(r"签订日期[:：]\s*(.+)"), "sign_date"),
        (re.compile(r"签订地点[:：]\s*(.+)"), "sign_location"),
        (re.compile(r"甲方[:：]\s*(.+)"), "party_a"),
        (re.compile(r"乙方[:：]\s*(.+)"), "party_b"),
        (re.compile(r"丙方[:：]\s*(.+)"), "party_c"),
    ]

    AMOUNT_PATTERNS = [
        re.compile(r"合同(?:总)?金额[:：]?\s*[￥¥$€£]?\s*([\d,]+(?:\.\d+)?)\s*(元|美元|欧元|英镑|CNY|USD|EUR|GBP)?"),
        re.compile(r"价款[:：]?\s*[￥¥$€£]?\s*([\d,]+(?:\.\d+)?)\s*(元|美元|欧元|英镑|CNY|USD|EUR|GBP)?"),
        re.compile(r"共计[￥¥$€£]?\s*([\d,]+(?:\.\d+)?)\s*(元|美元|欧元|英镑|CNY|USD|EUR|GBP)?"),
    ]

    def __init__(self, upload_dir: str = "./data/uploads"):
        self.upload_dir = Path(upload_dir)
        self.upload_dir.mkdir(parents=True, exist_ok=True)

    @classmethod
    def get_file_hash(cls, file_path: str) -> str:
        sha256_hash = hashlib.sha256()
        with open(file_path, "rb") as f:
            for byte_block in iter(lambda: f.read(4096), b""):
                sha256_hash.update(byte_block)
        return sha256_hash.hexdigest()

    @classmethod
    def detect_mime_type(cls, file_name: str) -> str:
        ext = Path(file_name).suffix.lower()
        return cls.SUPPORTED_MIME.get(ext, "application/octet-stream")

    def save_uploaded_file(self, file_content: bytes, original_name: str, uploader_id: Optional[int] = None) -> Tuple[str, int]:
        ext = Path(original_name).suffix.lower()
        if ext not in self.SUPPORTED_MIME:
            raise ValueError(f"不支持的文件类型: {ext}")

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_name = re.sub(r"[^\w\u4e00-\u9fff\.\-]", "_", original_name)
        new_name = f"{timestamp}_{uploader_id or 'anon'}_{safe_name}"
        full_path = self.upload_dir / new_name

        with open(full_path, "wb") as f:
            f.write(file_content)

        return str(full_path), len(file_content)

    def parse(self, file_path: str, original_name: Optional[str] = None) -> ParsedDocument:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")

        ext = path.suffix.lower()
        file_name = original_name or path.name
        file_size = path.stat().st_size
        file_hash = self.get_file_hash(file_path)
        mime_type = self.detect_mime_type(file_name)

        doc = ParsedDocument(
            file_name=file_name,
            file_path=str(path),
            file_size=file_size,
            file_hash=file_hash,
            mime_type=mime_type,
        )

        try:
            if ext == ".pdf":
                self._parse_pdf(path, doc)
            elif ext == ".docx":
                self._parse_docx(path, doc)
            elif ext == ".txt":
                self._parse_txt(path, doc)
            else:
                raise ValueError(f"不支持的文件类型: {ext}")

            self._extract_metadata(doc)
            self._split_clauses(doc)
            doc.word_count = len(re.findall(r"\S+", doc.full_text))

        except Exception as e:
            logger.exception(f"解析文档失败: {file_name}")
            doc.error_message = str(e)

        return doc

    def _parse_pdf(self, path: Path, doc: ParsedDocument) -> None:
        try:
            import pdfplumber
        except ImportError:
            logger.warning("pdfplumber未安装，使用pypdf作为后备")
            self._parse_pdf_pypdf(path, doc)
            return

        pages_text = []
        with pdfplumber.open(str(path)) as pdf:
            doc.page_count = len(pdf.pages)
            doc.metadata = pdf.metadata or {}
            char_offset = 0

            for page_idx, page in enumerate(pdf.pages, 1):
                text = page.extract_text() or ""
                pages_text.append(text)

        doc.full_text = "\n".join(pages_text)
        doc.title = doc.metadata.get("Title", "") or path.stem

    def _parse_pdf_pypdf(self, path: Path, doc: ParsedDocument) -> None:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        doc.page_count = len(reader.pages)
        doc.metadata = dict(reader.metadata) if reader.metadata else {}

        pages_text = []
        for page in reader.pages:
            text = page.extract_text() or ""
            pages_text.append(text)
        doc.full_text = "\n".join(pages_text)
        doc.title = doc.metadata.get("/Title", "") or path.stem

    def _parse_docx(self, path: Path, doc: ParsedDocument) -> None:
        from docx import Document
        document = Document(str(path))
        doc.page_count = 1
        doc.metadata = {
            "author": document.core_properties.author,
            "created": str(document.core_properties.created),
            "modified": str(document.core_properties.modified),
            "title": document.core_properties.title,
        }
        doc.title = document.core_properties.title or path.stem

        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        doc.full_text = "\n".join(paragraphs)

    def _parse_txt(self, path: Path, doc: ParsedDocument) -> None:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            doc.full_text = f.read()
        doc.page_count = max(1, doc.full_text.count("\n") // 50)
        doc.title = path.stem

    def _extract_metadata(self, doc: ParsedDocument) -> None:
        text = doc.full_text

        if not doc.title:
            match = re.search(r"^(.{5,100}合同.*?)$", text, re.MULTILINE)
            if match:
                doc.title = match.group(1).strip()

        for pattern, key in self.TITLE_PATTERNS:
            match = pattern.search(text)
            if match:
                value = match.group(1).strip()
                if key.endswith("_date"):
                    doc.key_dates[key] = self._parse_date(value)
                elif key.startswith("party_"):
                    doc.parties[key] = value
                elif key == "contract_name" and not doc.title:
                    doc.title = value
                doc.metadata[key] = value

        for pattern in self.AMOUNT_PATTERNS:
            match = pattern.search(text)
            if match:
                amount_str = match.group(1).replace(",", "")
                try:
                    doc.total_amount = float(amount_str)
                    if match.group(2):
                        currency_map = {
                            "元": "CNY", "CNY": "CNY", "人民币": "CNY",
                            "美元": "USD", "USD": "USD", "$": "USD",
                            "欧元": "EUR", "EUR": "EUR", "€": "EUR",
                            "英镑": "GBP", "GBP": "GBP", "£": "GBP",
                        }
                        doc.currency = currency_map.get(match.group(2), "CNY")
                except ValueError:
                    pass
                break

    @staticmethod
    def _parse_date(date_str: str) -> Optional[datetime]:
        date_formats = [
            "%Y年%m月%d日",
            "%Y-%m-%d",
            "%Y/%m/%d",
            "%Y.%m.%d",
            "%Y年%m月%d号",
        ]
        clean = re.sub(r"\s+", "", date_str)
        for fmt in date_formats:
            try:
                return datetime.strptime(clean, fmt)
            except ValueError:
                continue
        return None

    def _split_clauses(self, doc: ParsedDocument) -> None:
        text = doc.full_text
        if not text.strip():
            return

        split_positions = set()
        for pattern in self.CLAUSE_PATTERNS:
            for match in pattern.finditer(text):
                split_positions.add(match.start())

        if len(split_positions) < 3:
            split_positions = self._heuristic_split(text)

        sorted_positions = sorted(split_positions)

        current_char = 0
        for idx, pos in enumerate(sorted_positions):
            next_pos = sorted_positions[idx + 1] if idx + 1 < len(sorted_positions) else len(text)
            clause_text = text[pos:next_pos].strip()
            if not clause_text:
                continue

            title, number = self._extract_clause_title(clause_text)

            page_start = min(doc.page_count, pos // max(1, len(text) // doc.page_count) + 1)
            page_end = min(doc.page_count, next_pos // max(1, len(text) // doc.page_count) + 1)

            clause = ParsedClause(
                index=len(doc.clauses) + 1,
                title=title,
                number=number,
                text=clause_text,
                page_start=page_start,
                page_end=page_end,
                char_start=pos,
                char_end=next_pos,
            )
            doc.clauses.append(clause)

        if not doc.clauses:
            doc.clauses.append(ParsedClause(
                index=1,
                title="全文",
                text=text,
                char_start=0,
                char_end=len(text),
            ))

    @staticmethod
    def _heuristic_split(text: str) -> set:
        positions = set()
        lines = text.split("\n")
        char_offset = 0
        for line in lines:
            stripped = line.strip()
            if stripped and (
                (len(stripped) < 50 and (stripped.endswith("：") or stripped.endswith(":")))
                or (len(stripped) < 30 and re.match(r"^[\d一二三四五六七八九十]+[.、\)]", stripped))
            ):
                positions.add(char_offset)
            char_offset += len(line) + 1
        return positions

    @staticmethod
    def _extract_clause_title(clause_text: str) -> Tuple[Optional[str], Optional[str]]:
        first_line = clause_text.split("\n", 1)[0].strip()
        number = None

        num_match = re.match(r"^(第[一二三四五六七八九十百千万\d]+[章节条条款项]|\d+[\.、\)]|[一二三四五六七八九十]+[\.、\)]|Article\s+\d+|Section\s+\d+|Clause\s+\d+)\s*",
                            first_line, re.IGNORECASE)
        if num_match:
            number = num_match.group(1).strip()
            remaining = first_line[num_match.end():].strip()
            title = remaining if remaining and len(remaining) < 100 else None
        else:
            title = first_line if len(first_line) < 100 else None

        return title, number
