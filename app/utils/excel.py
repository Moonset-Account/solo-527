from openpyxl import Workbook, load_workbook
from io import BytesIO
from datetime import datetime
from app.config import ATTACHMENT_TYPE_NAMES, STATUS_NAMES, FINANCE_FIELDS

FIELD_NAMES_CN = {
    "order_no": "调拨单号",
    "batch_no": "批次号",
    "transfer_date": "调拨日期",
    "from_warehouse": "调出仓库",
    "to_warehouse": "调入仓库",
    "amount": "金额",
    "carrier": "承运商",
    "status": "状态",
    "created_by": "创建人",
    "created_at": "创建时间",
    "reviewed_by": "复核人",
    "reviewed_at": "复核时间",
    "review_comment": "复核意见",
}

def export_to_excel(rows: list[dict], field_names: list[str] = None, is_finance: bool = False) -> BytesIO:
    if is_finance:
        field_names = FINANCE_FIELDS
    
    if not field_names:
        field_names = list(FIELD_NAMES_CN.keys())
    
    wb = Workbook()
    ws = wb.active
    ws.title = "调拨单数据"
    
    headers = [FIELD_NAMES_CN.get(f, f) for f in field_names]
    ws.append(headers)
    
    for row in rows:
        data_row = []
        for f in field_names:
            val = row.get(f, "")
            if f == "status" and val in STATUS_NAMES:
                val = STATUS_NAMES[val]
            if isinstance(val, datetime):
                val = val.strftime("%Y-%m-%d %H:%M:%S")
            data_row.append(val)
        ws.append(data_row)
    
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column].width = adjusted_width
    
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def parse_excel(file_content: bytes) -> tuple[list[dict], list[str]]:
    errors = []
    data = []
    
    try:
        wb = load_workbook(filename=BytesIO(file_content), data_only=True)
        ws = wb.active
        
        headers = []
        for cell in ws[1]:
            headers.append(str(cell.value).strip() if cell.value else "")
        
        header_map = {}
        for idx, h in enumerate(headers):
            for key, cn in FIELD_NAMES_CN.items():
                if h == cn:
                    header_map[idx] = key
                    break
        
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            row_data = {}
            for col_idx, value in enumerate(row):
                if col_idx in header_map:
                    field = header_map[col_idx]
                    row_data[field] = value
            
            if row_data.get("order_no"):
                data.append(row_data)
            else:
                errors.append(f"第 {row_idx} 行缺少调拨单号")
    
    except Exception as e:
        errors.append(f"解析Excel失败: {str(e)}")
    
    return data, errors
