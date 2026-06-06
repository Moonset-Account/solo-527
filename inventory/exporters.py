import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from django.http import HttpResponse
from django.utils import timezone
from exhibitions.models import BorrowOrder, BorrowItem
from .models import InventoryItem


class ExportService:
    @staticmethod
    def export_borrow_order_to_excel(borrow_order):
        wb = Workbook()
        ws = wb.active
        ws.title = '借用单'

        header_font = Font(bold=True, size=12)
        header_fill = PatternFill(start_color='E0E0E0', end_color='E0E0E0', fill_type='solid')
        center_align = Alignment(horizontal='center', vertical='center')

        ws['A1'] = '博物馆临时展览物料借用单'
        ws['A1'].font = Font(bold=True, size=14)
        ws.merge_cells('A1:H1')
        ws['A1'].alignment = center_align

        ws['A3'] = '借单号:'
        ws['B3'] = borrow_order.order_no
        ws['C3'] = '展览名称:'
        ws['D3'] = borrow_order.exhibition.name
        ws['E3'] = '展厅:'
        ws['F3'] = borrow_order.hall.name
        ws['G3'] = '状态:'
        ws['H3'] = borrow_order.get_status_display_name()

        ws['A4'] = '申请人:'
        ws['B4'] = str(borrow_order.requester)
        ws['C4'] = '施工负责人:'
        ws['D4'] = str(borrow_order.construction_lead) if borrow_order.construction_lead else ''
        ws['E4'] = '借用类型:'
        ws['F4'] = borrow_order.get_type_display_name()
        ws['G4'] = '是否跨展厅:'
        ws['H4'] = '是' if borrow_order.is_cross_hall else '否'

        ws['A5'] = '预计领取日期:'
        ws['B5'] = str(borrow_order.expected_pickup_date)
        ws['C5'] = '预计归还日期:'
        ws['D5'] = str(borrow_order.expected_return_date)
        ws['E5'] = '实际领取时间:'
        ws['F5'] = str(borrow_order.actual_pickup_date) if borrow_order.actual_pickup_date else ''
        ws['G5'] = '实际归还时间:'
        ws['H5'] = str(borrow_order.actual_return_date) if borrow_order.actual_return_date else ''

        headers = ['序号', '物料编码', '物料名称', '物料类型', '规格型号', '数量', '已领取', '已归还', '状态']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=7, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align

        for idx, item in enumerate(borrow_order.items.all(), 1):
            row = 7 + idx
            ws.cell(row=row, column=1, value=idx).alignment = center_align
            ws.cell(row=row, column=2, value=item.material.code)
            ws.cell(row=row, column=3, value=item.material.name)
            ws.cell(row=row, column=4, value=item.material.get_type_display_name())
            ws.cell(row=row, column=5, value=item.material.specification)
            ws.cell(row=row, column=6, value=item.quantity).alignment = center_align
            ws.cell(row=row, column=7, value=item.picked_up_quantity).alignment = center_align
            ws.cell(row=row, column=8, value=item.returned_quantity).alignment = center_align
            ws.cell(row=row, column=9, value=item.get_status_display_name()).alignment = center_align

        for col in range(1, 10):
            ws.column_dimensions[chr(64 + col)].width = 15

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def export_unreturned_list_to_excel(date=None):
        from django.utils import timezone
        if date is None:
            date = timezone.now().date()

        wb = Workbook()
        ws = wb.active
        ws.title = '未归还清单'

        header_font = Font(bold=True, size=12)
        header_fill = PatternFill(start_color='FFE0E0', end_color='FFE0E0', fill_type='solid')

        ws['A1'] = f'未归还物料清单 - {date}'
        ws['A1'].font = Font(bold=True, size=14)
        ws.merge_cells('A1:I1')

        headers = ['借单号', '展览名称', '展厅', '物料编码', '物料名称', '应还数量', '已还数量', '未还数量', '逾期天数']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill

        from exhibitions.models import BorrowItem
        unreturned_items = BorrowItem.objects.filter(
            status__in=['picked_up', 'pending'],
            borrow_order__expected_return_date__lt=date
        ).select_related('borrow_order', 'borrow_order__exhibition', 'borrow_order__hall', 'material')

        for idx, item in enumerate(unreturned_items, 1):
            row = 3 + idx
            overdue_days = (date - item.borrow_order.expected_return_date).days
            unreturned_qty = item.quantity - item.returned_quantity

            ws.cell(row=row, column=1, value=item.borrow_order.order_no)
            ws.cell(row=row, column=2, value=item.borrow_order.exhibition.name)
            ws.cell(row=row, column=3, value=item.borrow_order.hall.name)
            ws.cell(row=row, column=4, value=item.material.code)
            ws.cell(row=row, column=5, value=item.material.name)
            ws.cell(row=row, column=6, value=item.quantity)
            ws.cell(row=row, column=7, value=item.returned_quantity)
            ws.cell(row=row, column=8, value=unreturned_qty)
            ws.cell(row=row, column=9, value=overdue_days)

        for col in range(1, 10):
            ws.column_dimensions[chr(64 + col)].width = 18

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def export_inventory_to_excel():
        wb = Workbook()
        ws = wb.active
        ws.title = '库存清单'

        header_font = Font(bold=True, size=12)
        header_fill = PatternFill(start_color='E0E0FF', end_color='E0E0FF', fill_type='solid')

        ws['A1'] = '库存清单'
        ws['A1'].font = Font(bold=True, size=14)
        ws.merge_cells('A1:H1')

        headers = ['物料编码', '物料名称', '物料类型', '序列号', '仓库', '状态', '具体位置', '备注']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill

        items = InventoryItem.objects.select_related('material', 'warehouse').all()
        for idx, item in enumerate(items, 1):
            row = 3 + idx
            ws.cell(row=row, column=1, value=item.material.code)
            ws.cell(row=row, column=2, value=item.material.name)
            ws.cell(row=row, column=3, value=item.material.get_type_display_name())
            ws.cell(row=row, column=4, value=item.serial_number)
            ws.cell(row=row, column=5, value=item.warehouse.name if item.warehouse else '')
            ws.cell(row=row, column=6, value=item.get_status_display_name())
            ws.cell(row=row, column=7, value=item.location_detail)
            ws.cell(row=row, column=8, value=item.remarks)

        for col in range(1, 9):
            ws.column_dimensions[chr(64 + col)].width = 18

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output
