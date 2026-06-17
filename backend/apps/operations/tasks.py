from celery import shared_task
from django.core.mail import send_mail
from django.conf import settings
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import os
import datetime
import logging
from .models import ExportRecord
from apps.tickets.models import Ticket
from apps.accounts.models import User

logger = logging.getLogger(__name__)


@shared_task(bind=True)
def export_tickets_to_excel(self, export_record_id):
    try:
        export_record = ExportRecord.objects.get(id=export_record_id)
        export_record.status = 'processing'
        export_record.save()
        
        filters = export_record.filters or {}
        
        queryset = Ticket.objects.all()
        
        if 'status' in filters:
            queryset = queryset.filter(status=filters['status'])
        if 'priority' in filters:
            queryset = queryset.filter(priority=filters['priority'])
        if 'type' in filters:
            queryset = queryset.filter(type=filters['type'])
        if 'assignee_id' in filters:
            queryset = queryset.filter(assignee_id=filters['assignee_id'])
        if 'creator_id' in filters:
            queryset = queryset.filter(creator_id=filters['creator_id'])
        if 'start_date' in filters:
            queryset = queryset.filter(created_at__gte=filters['start_date'])
        if 'end_date' in filters:
            queryset = queryset.filter(created_at__lte=filters['end_date'])
        
        export_record.record_count = queryset.count()
        export_record.save()
        
        wb = Workbook()
        ws = wb.active
        ws.title = '工单数据'
        
        headers = [
            '工单编号', '标题', '类型', '状态', '优先级', 
            '客户姓名', '客户电话', '订单号', '商品名称',
            '处理人', '创建人', '创建时间', 'SLA截止时间',
            '首次响应时间', '解决时间', '处理结果', '备注'
        ]
        
        header_fill = PatternFill(start_color='1E40AF', end_color='1E40AF', fill_type='solid')
        header_font = Font(color='FFFFFF', bold=True)
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        for row, ticket in enumerate(queryset, 2):
            ws.cell(row=row, column=1, value=ticket.ticket_no)
            ws.cell(row=row, column=2, value=ticket.title)
            ws.cell(row=row, column=3, value=ticket.get_type_display())
            ws.cell(row=row, column=4, value=ticket.get_status_display())
            ws.cell(row=row, column=5, value=ticket.get_priority_display())
            ws.cell(row=row, column=6, value=ticket.customer_name)
            ws.cell(row=row, column=7, value=ticket.customer_phone)
            ws.cell(row=row, column=8, value=ticket.order_no or '')
            ws.cell(row=row, column=9, value=ticket.product_name or '')
            ws.cell(row=row, column=10, value=ticket.assignee.name if ticket.assignee else '')
            ws.cell(row=row, column=11, value=ticket.creator.name if ticket.creator else '')
            ws.cell(row=row, column=12, value=ticket.created_at.strftime('%Y-%m-%d %H:%M:%S'))
            ws.cell(row=row, column=13, value=ticket.sla_deadline.strftime('%Y-%m-%d %H:%M:%S'))
            ws.cell(row=row, column=14, value=ticket.first_response_at.strftime('%Y-%m-%d %H:%M:%S') if ticket.first_response_at else '')
            ws.cell(row=row, column=15, value=ticket.resolved_at.strftime('%Y-%m-%d %H:%M:%S') if ticket.resolved_at else '')
            ws.cell(row=row, column=16, value=ticket.resolution or '')
        
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + col)].width = 15
        
        exports_dir = os.path.join(settings.MEDIA_ROOT, 'exports')
        os.makedirs(exports_dir, exist_ok=True)
        
        filename = f'tickets_export_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx'
        filepath = os.path.join(exports_dir, filename)
        
        wb.save(filepath)
        
        export_record.file_name = filename
        export_record.file_path = f'exports/{filename}'
        export_record.status = 'completed'
        export_record.completed_at = datetime.datetime.now()
        export_record.save()
        
        if export_record.created_by and export_record.created_by.email:
            send_mail(
                subject='【导出完成】工单数据导出成功',
                message=f'''
                您的工单数据导出已完成。
                
                导出记录ID: {export_record.id}
                文件名: {filename}
                记录数: {export_record.record_count}
                
                请登录系统下载文件。
                ''',
                from_email=settings.DEFAULT_FROM_EMAIL,
                recipient_list=[export_record.created_by.email],
                fail_silently=True
            )
        
        logger.info(f'Export {export_record_id} completed successfully')
        return f'Exported {export_record.record_count} records'
        
    except Exception as e:
        logger.error(f'Export {export_record_id} failed: {str(e)}')
        export_record.status = 'failed'
        export_record.error_message = str(e)
        export_record.save()
        raise self.retry(exc=e, countdown=60, max_retries=3)


@shared_task
def send_improvement_due_reminder():
    from .models import ImprovementAction
    
    today = datetime.date.today()
    threshold = today + datetime.timedelta(days=3)
    
    upcoming = ImprovementAction.objects.filter(
        status__in=['pending', 'in_progress'],
        due_date__lte=threshold,
        due_date__gte=today
    )
    
    overdue = ImprovementAction.objects.filter(
        status__in=['pending', 'in_progress'],
        due_date__lt=today
    )
    
    for action in upcoming:
        if action.assignee and action.assignee.email:
            send_mail(
                subject=f'【改进项即将到期】{action.title}',
                message=f'''
                您负责的改进项即将到期：
                
                标题: {action.title}
                状态: {action.get_status_display()}
                优先级: {action.get_priority_display()}
                截止日期: {action.due_date}
                当前进度: {action.progress}%
                
                请及时处理！
                ''',
                from_email=settings.DEFAULT_FROM_EMAIL,
                recipient_list=[action.assignee.email],
                fail_silently=True
            )
    
    for action in overdue:
        if action.assignee and action.assignee.email:
            send_mail(
                subject=f'【改进项已逾期】{action.title}',
                message=f'''
                您负责的改进项已逾期：
                
                标题: {action.title}
                状态: {action.get_status_display()}
                优先级: {action.get_priority_display()}
                截止日期: {action.due_date}
                当前进度: {action.progress}%
                
                请尽快处理！
                ''',
                from_email=settings.DEFAULT_FROM_EMAIL,
                recipient_list=[action.assignee.email],
                fail_silently=True
            )
    
    return f'Sent {upcoming.count() + overdue.count()} reminders'
