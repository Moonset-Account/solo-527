from rest_framework import viewsets, status, permissions
from rest_framework.decorators import action
from rest_framework.response import Response
from django.db import transaction, models
from django.db.models import Q, Count, Avg, F
from django.utils import timezone
from django_filters.rest_framework import DjangoFilterBackend, FilterSet, DateFilter, NumberFilter
from rest_framework.filters import SearchFilter, OrderingFilter
from datetime import timedelta

from .models import RepairRequest, RepairPhoto, RepairProgress, RepairComment, RepairStatus
from .serializers import (
    RepairRequestSerializer, RepairRequestCreateSerializer, RepairPhotoSerializer,
    RepairProgressSerializer, RepairCommentSerializer, RepairAssignSerializer,
    RepairStatusUpdateSerializer,
)
from apps.users.models import User, Role
from apps.audit.utils import log_audit
from apps.audit.models import OperationType


class RepairFilterSet(FilterSet):
    created_from = DateFilter(field_name='created_at', lookup_expr='date__gte')
    created_to = DateFilter(field_name='created_at', lookup_expr='date__lte')
    completed_from = DateFilter(field_name='completed_at', lookup_expr='date__gte')
    completed_to = DateFilter(field_name='completed_at', lookup_expr='date__lte')
    min_processing_hours = NumberFilter(method='filter_min_processing')
    max_processing_hours = NumberFilter(method='filter_max_processing')

    class Meta:
        model = RepairRequest
        fields = ['status', 'repair_type', 'priority', 'dorm_building', 'applicant', 'assignee']

    def filter_min_processing(self, queryset, name, value):
        return queryset.filter(
            completed_at__isnull=False
        ).extra(
            where=['EXTRACT(EPOCH FROM (completed_at - created_at)) / 3600 >= %s'],
            params=[value]
        )

    def filter_max_processing(self, queryset, name, value):
        return queryset.filter(
            completed_at__isnull=False
        ).extra(
            where=['EXTRACT(EPOCH FROM (completed_at - created_at)) / 3600 <= %s'],
            params=[value]
        )


class CanManageRepair(permissions.BasePermission):
    def has_permission(self, request, view):
        if not request.user.is_authenticated:
            return False
        if view.action in ['list', 'retrieve']:
            return True
        return request.user.role in [Role.ADMIN, Role.DORM_MANAGER, Role.MAINTENANCE]

    def has_object_permission(self, request, view, obj):
        if request.user.role in [Role.ADMIN, Role.DORM_MANAGER]:
            return True
        if request.user.role == Role.MAINTENANCE and obj.assignee == request.user:
            return True
        return obj.applicant == request.user


class RepairRequestViewSet(viewsets.ModelViewSet):
    queryset = RepairRequest.objects.all()
    filter_backends = [DjangoFilterBackend, SearchFilter, OrderingFilter]
    filterset_class = RepairFilterSet
    search_fields = ['title', 'description', 'contact_name', 'contact_phone']
    ordering_fields = ['created_at', 'updated_at', 'completed_at', 'priority']
    permission_classes = [CanManageRepair]

    def get_serializer_class(self):
        if self.action == 'create':
            return RepairRequestCreateSerializer
        return RepairRequestSerializer

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user
        if user.role == Role.STUDENT:
            return qs.filter(applicant=user)
        if user.role == Role.MAINTENANCE:
            return qs.filter(Q(assignee=user) | Q(status=RepairStatus.PENDING))
        if user.role == Role.DORM_MANAGER:
            return qs.filter(dorm_building=user.dorm_building) if user.dorm_building else qs
        return qs

    def perform_create(self, serializer):
        with transaction.atomic():
            repair = serializer.save()
            log_audit(
                self.request.user, OperationType.CREATE, '报修管理',
                f'提交报修: {repair.title}',
                new_data={'id': repair.id, 'title': repair.title},
                content_object=repair,
                ip_address=self._get_ip()
            )

    @action(detail=True, methods=['post'])
    def assign(self, request, pk=None):
        if request.user.role not in [Role.ADMIN, Role.DORM_MANAGER]:
            return Response({'detail': '无权限'}, status=status.HTTP_403_FORBIDDEN)
        repair = self.get_object()
        serializer = RepairAssignSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        assignee_id = serializer.validated_data['assignee_id']
        try:
            assignee = User.objects.get(id=assignee_id, role__in=[Role.MAINTENANCE, Role.DORM_MANAGER, Role.ADMIN])
        except User.DoesNotExist:
            return Response({'detail': '处理人不存在或角色不匹配'}, status=status.HTTP_400_BAD_REQUEST)

        with transaction.atomic():
            old_status = repair.status
            repair.assignee = assignee
            repair.status = RepairStatus.ASSIGNED
            repair.save()
            RepairProgress.objects.create(
                repair_request=repair,
                status=RepairStatus.ASSIGNED,
                remark=serializer.validated_data.get('remark', f'分配给 {assignee.real_name or assignee.username}'),
                operator=request.user
            )
            log_audit(
                request.user, OperationType.ASSIGN, '报修管理',
                f'分配报修 #{repair.id} 给 {assignee}',
                content_object=repair,
                ip_address=self._get_ip()
            )
        return Response(RepairRequestSerializer(repair).data)

    @action(detail=True, methods=['post'])
    def update_status(self, request, pk=None):
        repair = self.get_object()
        if not self._can_update_status(request.user, repair):
            return Response({'detail': '无权限'}, status=status.HTTP_403_FORBIDDEN)
        serializer = RepairStatusUpdateSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        new_status = serializer.validated_data['status']
        remark = serializer.validated_data.get('remark', '')

        with transaction.atomic():
            old_status = repair.status
            repair.status = new_status
            if new_status == RepairStatus.COMPLETED:
                repair.completed_at = timezone.now()
            repair.save()
            RepairProgress.objects.create(
                repair_request=repair,
                status=new_status,
                remark=remark,
                operator=request.user
            )
            log_audit(
                request.user,
                OperationType.COMPLETE if new_status == RepairStatus.COMPLETED else OperationType.UPDATE,
                '报修管理',
                f'报修 #{repair.id} 状态更新为 {repair.get_status_display()}',
                old_data={'status': old_status},
                new_data={'status': new_status},
                content_object=repair,
                ip_address=self._get_ip()
            )
        return Response(RepairRequestSerializer(repair).data)

    @action(detail=True, methods=['post'])
    def comment(self, request, pk=None):
        repair = self.get_object()
        serializer = RepairCommentSerializer(data=request.data, context={'request': request})
        serializer.is_valid(raise_exception=True)
        comment = serializer.save(user=request.user, repair_request=repair)
        return Response(RepairCommentSerializer(comment).data, status=status.HTTP_201_CREATED)

    @action(detail=False, methods=['get'])
    def statistics(self, request):
        if request.user.role not in [Role.ADMIN, Role.DORM_MANAGER]:
            return Response({'detail': '无权限'}, status=status.HTTP_403_FORBIDDEN)
        qs = self.get_queryset()
        days = int(request.GET.get('days', 30))
        start_date = timezone.now() - timedelta(days=days)
        qs = qs.filter(created_at__gte=start_date)

        by_status = dict(qs.values_list('status').annotate(count=Count('id')))
        by_type = dict(qs.values_list('repair_type').annotate(count=Count('id')))
        completed = qs.filter(status=RepairStatus.COMPLETED)
        avg_time = completed.aggregate(
            avg=Avg(F('completed_at') - F('created_at'))
        )['avg']
        avg_hours = avg_time.total_seconds() / 3600 if avg_time else None

        return Response({
            'total': qs.count(),
            'by_status': by_status,
            'by_type': by_type,
            'avg_processing_hours': round(avg_hours, 2) if avg_hours else None,
            'completed_count': completed.count(),
        })

    @action(detail=False, methods=['get'])
    def export(self, request):
        if request.user.role not in [Role.ADMIN, Role.DORM_MANAGER]:
            return Response({'detail': '无权限'}, status=status.HTTP_403_FORBIDDEN)
        from django.http import HttpResponse
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment

        qs = self.filter_queryset(self.get_queryset())

        wb = Workbook()
        ws = wb.active
        ws.title = '报修记录'

        headers = ['ID', '标题', '类型', '优先级', '状态', '宿舍楼', '宿舍号', '联系人', '联系电话',
                   '申请人', '处理人', '创建时间', '完成时间', '处理时长(小时)']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color='DDDDDD', end_color='DDDDDD', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')

        for row, repair in enumerate(qs, 2):
            processing = None
            if repair.completed_at:
                processing = round((repair.completed_at - repair.created_at).total_seconds() / 3600, 2)
            ws.cell(row=row, column=1, value=repair.id)
            ws.cell(row=row, column=2, value=repair.title)
            ws.cell(row=row, column=3, value=repair.get_repair_type_display())
            ws.cell(row=row, column=4, value=repair.get_priority_display())
            ws.cell(row=row, column=5, value=repair.get_status_display())
            ws.cell(row=row, column=6, value=repair.dorm_building)
            ws.cell(row=row, column=7, value=repair.dorm_room)
            ws.cell(row=row, column=8, value=repair.contact_name)
            ws.cell(row=row, column=9, value=repair.contact_phone)
            ws.cell(row=row, column=10, value=str(repair.applicant))
            ws.cell(row=row, column=11, value=str(repair.assignee) if repair.assignee else '')
            ws.cell(row=row, column=12, value=repair.created_at.strftime('%Y-%m-%d %H:%M:%S'))
            ws.cell(row=row, column=13, value=repair.completed_at.strftime('%Y-%m-%d %H:%M:%S') if repair.completed_at else '')
            ws.cell(row=row, column=14, value=processing)

        for col in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + col)].width = 18

        log_audit(
            request.user, OperationType.EXPORT, '报修管理',
            f'导出报修记录，共 {qs.count()} 条',
            ip_address=self._get_ip()
        )

        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename=repairs_{timezone.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        wb.save(response)
        return response

    def _can_update_status(self, user, repair):
        if user.role in [Role.ADMIN, Role.DORM_MANAGER]:
            return True
        if user.role == Role.MAINTENANCE and repair.assignee == user:
            return True
        return False

    def _get_ip(self):
        return self.request.META.get('REMOTE_ADDR')
