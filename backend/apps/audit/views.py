from rest_framework import viewsets, permissions, status
from rest_framework.decorators import action
from rest_framework.response import Response
from django_filters.rest_framework import DjangoFilterBackend, FilterSet, DateFilter
from rest_framework.filters import SearchFilter, OrderingFilter
from django.utils import timezone
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from .models import AuditLog
from .serializers import AuditLogSerializer
from apps.users.models import Role


class AuditLogFilterSet(FilterSet):
    created_from = DateFilter(field_name='created_at', lookup_expr='date__gte')
    created_to = DateFilter(field_name='created_at', lookup_expr='date__lte')

    class Meta:
        model = AuditLog
        fields = ['user', 'operation', 'module']


class AuditLogViewSet(viewsets.ReadOnlyModelViewSet):
    queryset = AuditLog.objects.all()
    serializer_class = AuditLogSerializer
    filter_backends = [DjangoFilterBackend, SearchFilter, OrderingFilter]
    filterset_class = AuditLogFilterSet
    search_fields = ['username', 'description', 'module']
    ordering_fields = ['created_at']

    def get_permissions(self):
        return [permissions.IsAuthenticated()]

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user
        if user.role not in [Role.ADMIN]:
            if user.role == Role.DORM_MANAGER:
                return qs.filter(user=user) | qs.filter(module__in=['报修管理', '自习室管理'])
            return qs.filter(user=user)
        return qs

    @action(detail=False, methods=['get'])
    def export(self, request):
        if request.user.role not in [Role.ADMIN]:
            return Response({'detail': '无权限'}, status=status.HTTP_403_FORBIDDEN)

        qs = self.filter_queryset(self.get_queryset())
        wb = Workbook()
        ws = wb.active
        ws.title = '操作日志'

        headers = ['ID', '用户名', '操作类型', '模块', '描述', 'IP地址', '操作时间']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color='DDDDDD', end_color='DDDDDD', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')

        for row, log in enumerate(qs, 2):
            ws.cell(row=row, column=1, value=log.id)
            ws.cell(row=row, column=2, value=log.username)
            ws.cell(row=row, column=3, value=log.get_operation_display())
            ws.cell(row=row, column=4, value=log.module)
            ws.cell(row=row, column=5, value=log.description)
            ws.cell(row=row, column=6, value=log.ip_address or '')
            ws.cell(row=row, column=7, value=log.created_at.strftime('%Y-%m-%d %H:%M:%S'))

        for col in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + col)].width = 20

        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename=audit_logs_{timezone.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        wb.save(response)
        return response
