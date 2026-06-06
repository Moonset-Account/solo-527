from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from django.http import HttpResponse
from datetime import datetime
import io


class ExcelExporter:
    def __init__(self, filename_prefix='export'):
        self.filename_prefix = filename_prefix
        self.wb = Workbook()
        self.ws = self.wb.active
        
    def set_headers(self, headers):
        self.ws.append(headers)
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        
        for cell in self.ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
    
    def add_rows(self, rows):
        for row in rows:
            self.ws.append(row)
    
    def adjust_column_width(self):
        for column in self.ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            self.ws.column_dimensions[column_letter].width = adjusted_width
    
    def get_response(self):
        self.adjust_column_width()
        buffer = io.BytesIO()
        self.wb.save(buffer)
        buffer.seek(0)
        
        filename = f'{self.filename_prefix}_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response


def export_borrow_records(queryset):
    exporter = ExcelExporter('借阅记录')
    headers = ['ID', '家庭', '绘本', '条码号', '借阅人', '借阅状态', '预约时间', '取书时间', '应还日期', '归还时间', '逾期罚款', '续借次数']
    exporter.set_headers(headers)
    
    rows = []
    for record in queryset:
        rows.append([
            record.id,
            record.family.name if record.family else '',
            record.book.title if record.book else '',
            record.book_copy.barcode if record.book_copy else '',
            record.borrower.username if record.borrower else '',
            record.get_status_display(),
            record.reserved_at.strftime('%Y-%m-%d %H:%M') if record.reserved_at else '',
            record.picked_up_at.strftime('%Y-%m-%d %H:%M') if record.picked_up_at else '',
            record.due_date.strftime('%Y-%m-%d') if record.due_date else '',
            record.returned_at.strftime('%Y-%m-%d %H:%M') if record.returned_at else '',
            float(record.overdue_fine),
            record.renew_count,
        ])
    exporter.add_rows(rows)
    return exporter.get_response()


def export_repair_records(queryset):
    exporter = ExcelExporter('修复记录')
    headers = ['ID', '绘本', '条码号', '上报人', '破损类型', '修复状态', '优先级', '预估费用', '实际费用', '上报时间', '完成时间', '负责人']
    exporter.set_headers(headers)
    
    rows = []
    for record in queryset:
        rows.append([
            record.id,
            record.book_copy.book.title if record.book_copy and record.book_copy.book else '',
            record.book_copy.barcode if record.book_copy else '',
            record.reported_by.username if record.reported_by else '',
            record.get_damage_type_display(),
            record.get_status_display(),
            record.get_priority_display(),
            float(record.estimated_cost),
            float(record.actual_cost),
            record.reported_at.strftime('%Y-%m-%d %H:%M') if record.reported_at else '',
            record.completed_at.strftime('%Y-%m-%d %H:%M') if record.completed_at else '',
            record.assigned_to.username if record.assigned_to else '',
        ])
    exporter.add_rows(rows)
    return exporter.get_response()


def export_activity_registrations(queryset):
    exporter = ExcelExporter('活动报名')
    headers = ['ID', '活动', '家庭', '儿童', '报名人', '报名状态', '候补位置', '报名时间', '押金支付', '签到时间']
    exporter.set_headers(headers)
    
    rows = []
    for reg in queryset:
        rows.append([
            reg.id,
            reg.activity.title if reg.activity else '',
            reg.family.name if reg.family else '',
            reg.child.name if reg.child else '',
            reg.registered_by.username if reg.registered_by else '',
            reg.get_status_display(),
            reg.waitlist_position if reg.waitlist_position else '',
            reg.registered_at.strftime('%Y-%m-%d %H:%M') if reg.registered_at else '',
            '是' if reg.deposit_paid else '否',
            reg.attended_at.strftime('%Y-%m-%d %H:%M') if reg.attended_at else '',
        ])
    exporter.add_rows(rows)
    return exporter.get_response()


def export_deposit_transactions(queryset):
    exporter = ExcelExporter('押金交易')
    headers = ['ID', '家庭', '交易类型', '金额', '可用余额', '冻结金额', '交易说明', '操作人', '交易时间', '是否已确认']
    exporter.set_headers(headers)
    
    rows = []
    for trans in queryset:
        rows.append([
            trans.id,
            trans.deposit.family.name if trans.deposit and trans.deposit.family else '',
            trans.get_transaction_type_display(),
            float(trans.amount),
            float(trans.balance_after),
            float(trans.frozen_after),
            trans.description,
            trans.operator.username if trans.operator else '',
            trans.created_at.strftime('%Y-%m-%d %H:%M') if trans.created_at else '',
            '是' if trans.confirmed_at else '否',
        ])
    exporter.add_rows(rows)
    return exporter.get_response()
