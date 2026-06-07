import pandas as pd
from io import BytesIO
from datetime import datetime
from typing import List, Dict, Any
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from sqlalchemy.orm import Session
from backend.db.models import Member, Checkin, Booking, Coach, Course, Store

class ExportService:
    
    @staticmethod
    def export_cohort_analysis(cohort_data: List[Dict]) -> BytesIO:
        output = BytesIO()
        
        wb = Workbook()
        ws = wb.active
        ws.title = "留存Cohort分析"
        
        headers = ["入组月份", "样本量", "M0", "M1", "M2", "M3", "M4", "M5", "M6"]
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        
        for row, cohort in enumerate(cohort_data, 2):
            ws.cell(row=row, column=1, value=cohort.get("cohort_month", ""))
            ws.cell(row=row, column=2, value=cohort.get("cohort_size", 0))
            
            retention = cohort.get("retention", [])
            for col, rate in enumerate(retention, 3):
                cell = ws.cell(row=row, column=col, value=f"{rate * 100:.1f}%" if rate else "-")
                cell.alignment = Alignment(horizontal="center")
                
                if rate and rate < 0.3:
                    cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                elif rate and rate > 0.6:
                    cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + col)].width = 14
        
        wb.save(output)
        output.seek(0)
        return output
    
    @staticmethod
    def export_churn_warning(members: List[Dict]) -> BytesIO:
        output = BytesIO()
        
        wb = Workbook()
        ws = wb.active
        ws.title = "流失预警列表"
        
        headers = ["会员ID", "姓名", "会员类型", "所属门店", "入会日期", 
                   "最后训练日期", "连续未训练天数", "历史周均频次", "流失风险等级"]
        
        header_fill = PatternFill(start_color="E74C3C", end_color="E74C3C", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
        
        for row, member in enumerate(members, 2):
            ws.cell(row=row, column=1, value=member.get("member_id", ""))
            ws.cell(row=row, column=2, value=member.get("name", ""))
            ws.cell(row=row, column=3, value=member.get("member_type", ""))
            ws.cell(row=row, column=4, value=member.get("store", ""))
            ws.cell(row=row, column=5, value=member.get("join_date", ""))
            ws.cell(row=row, column=6, value=member.get("last_checkin", ""))
            ws.cell(row=row, column=7, value=member.get("days_inactive", 0))
            ws.cell(row=row, column=8, value=f"{member.get('avg_weekly_freq', 0):.1f}")
            
            risk = member.get("risk_level", "")
            cell = ws.cell(row=row, column=9, value=risk)
            if risk == "高风险":
                cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            elif risk == "中风险":
                cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
        
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + col)].width = 16
        
        wb.save(output)
        output.seek(0)
        return output
    
    @staticmethod
    def export_coach_load(coach_data: List[Dict]) -> BytesIO:
        output = BytesIO()
        
        wb = Workbook()
        ws = wb.active
        ws.title = "教练负载分析"
        
        headers = ["教练姓名", "级别", "所属门店", "周均课时", "学员数", "私教转化率", "学员留存率"]
        
        header_fill = PatternFill(start_color="2ECC71", end_color="2ECC71", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
        
        for row, coach in enumerate(coach_data, 2):
            ws.cell(row=row, column=1, value=coach.get("name", ""))
            ws.cell(row=row, column=2, value=coach.get("level", ""))
            ws.cell(row=row, column=3, value=coach.get("store", ""))
            ws.cell(row=row, column=4, value=coach.get("weekly_hours", 0))
            ws.cell(row=row, column=5, value=coach.get("student_count", 0))
            ws.cell(row=row, column=6, value=f"{coach.get('pt_conversion', 0) * 100:.1f}%")
            ws.cell(row=row, column=7, value=f"{coach.get('retention_rate', 0) * 100:.1f}%")
        
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + col)].width = 16
        
        wb.save(output)
        output.seek(0)
        return output
