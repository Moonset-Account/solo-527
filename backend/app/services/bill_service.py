from typing import Optional, List
from datetime import datetime, date
from sqlalchemy.orm import Session
from sqlalchemy import and_, or_, func
import uuid
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, Alignment

from app.models.bill import Bill, BillPayment
from app.models.lease_ext import Lease
from app.schemas.bill import (
    BillCreate, BillUpdate, BillQuery,
    BillPaymentCreate, BillGenerateRequest,
    CollectionProgressResponse, CollectionProgressItem,
)
from app.schemas.common import PageResult


class BillService:
    @staticmethod
    def get(db: Session, bill_id: int) -> Optional[Bill]:
        return db.query(Bill).filter(Bill.id == bill_id, Bill.is_deleted == False).first()

    @staticmethod
    def get_by_no(db: Session, bill_no: str) -> Optional[Bill]:
        return db.query(Bill).filter(Bill.bill_no == bill_no, Bill.is_deleted == False).first()

    @staticmethod
    def list(db: Session, query: BillQuery) -> PageResult:
        q = db.query(Bill).filter(Bill.is_deleted == False)

        if query.keyword:
            keyword = f"%{query.keyword}%"
            q = q.join(Lease, Bill.lease_id == Lease.id).filter(
                or_(Bill.bill_no.like(keyword), Lease.lease_no.like(keyword))
            )

        if query.bill_type:
            q = q.filter(Bill.bill_type == query.bill_type)

        if query.status:
            q = q.filter(Bill.status == query.status)

        if query.lease_id:
            q = q.filter(Bill.lease_id == query.lease_id)

        if query.bill_date_from:
            q = q.filter(Bill.bill_date >= query.bill_date_from)

        if query.bill_date_to:
            q = q.filter(Bill.bill_date <= query.bill_date_to)

        if query.due_date_from:
            q = q.filter(Bill.due_date >= query.due_date_from)

        if query.due_date_to:
            q = q.filter(Bill.due_date <= query.due_date_to)

        total = q.count()
        items = q.order_by(Bill.id.desc()).offset((query.page - 1) * query.page_size).limit(query.page_size).all()

        return PageResult(total=total, page=query.page, page_size=query.page_size, items=items)

    @staticmethod
    def generate_bill_no() -> str:
        date_str = datetime.now().strftime("%Y%m%d")
        unique_id = str(uuid.uuid4().hex)[:8].upper()
        return f"ZD{date_str}{unique_id}"

    @staticmethod
    def create(db: Session, data: BillCreate, created_by: Optional[int] = None) -> Bill:
        bill_no = BillService.generate_bill_no()
        db_obj = Bill(**data.model_dump(), bill_no=bill_no, created_by=created_by)
        db.add(db_obj)
        db.commit()
        db.refresh(db_obj)
        return db_obj

    @staticmethod
    def update(db: Session, bill_id: int, data: BillUpdate, updated_by: Optional[int] = None) -> Optional[Bill]:
        db_obj = BillService.get(db, bill_id)
        if not db_obj:
            return None

        update_data = data.model_dump(exclude_unset=True)
        for field, value in update_data.items():
            setattr(db_obj, field, value)

        db_obj.updated_by = updated_by
        db_obj.updated_at = datetime.utcnow()

        db.commit()
        db.refresh(db_obj)
        return db_obj

    @staticmethod
    def delete(db: Session, bill_id: int, updated_by: Optional[int] = None) -> bool:
        db_obj = BillService.get(db, bill_id)
        if not db_obj:
            return False
        db_obj.is_deleted = True
        db_obj.updated_by = updated_by
        db_obj.updated_at = datetime.utcnow()
        db.commit()
        return True

    @staticmethod
    def batch_generate(db: Session, req: BillGenerateRequest, created_by: Optional[int] = None) -> int:
        lease_query = db.query(Lease).filter(Lease.is_deleted == False, Lease.status == "active")

        if req.lease_ids:
            lease_query = lease_query.filter(Lease.id.in_(req.lease_ids))

        leases = lease_query.all()
        count = 0

        for lease in leases:
            if req.bill_type == "rent":
                amount = float(lease.rent_amount)
            elif req.bill_type == "deposit":
                amount = float(lease.deposit_amount)
            else:
                amount = float(lease.rent_amount)

            existing = db.query(Bill).filter(
                Bill.lease_id == lease.id,
                Bill.bill_type == req.bill_type,
                Bill.bill_period == req.bill_period,
                Bill.is_deleted == False,
            ).first()

            if not existing:
                bill_no = BillService.generate_bill_no()
                bill = Bill(
                    bill_no=bill_no,
                    lease_id=lease.id,
                    bill_type=req.bill_type,
                    bill_period=req.bill_period,
                    bill_date=req.bill_date,
                    due_date=req.due_date,
                    amount=amount,
                    paid_amount=0,
                    status="pending",
                    created_by=created_by,
                )
                db.add(bill)
                count += 1

        db.commit()
        return count

    @staticmethod
    def get_collection_progress(db: Session) -> CollectionProgressResponse:
        results = (
            db.query(
                Bill.bill_period,
                func.sum(Bill.amount).label("total_amount"),
                func.sum(Bill.paid_amount).label("paid_amount"),
                func.count(Bill.id).label("bill_count"),
                func.sum(func.case((Bill.status == "paid", 1), else_=0)).label("paid_count"),
            )
            .filter(Bill.is_deleted == False, Bill.bill_type == "rent")
            .group_by(Bill.bill_period)
            .order_by(Bill.bill_period.desc())
            .limit(12)
            .all()
        )

        items = []
        total_amount = 0
        total_paid = 0

        for r in results:
            total = float(r.total_amount or 0)
            paid = float(r.paid_amount or 0)
            unpaid = total - paid
            rate = (paid / total * 100) if total > 0 else 0

            items.append(
                CollectionProgressItem(
                    period=r.bill_period or "",
                    total_amount=total,
                    paid_amount=paid,
                    unpaid_amount=unpaid,
                    collection_rate=round(rate, 2),
                    bill_count=r.bill_count or 0,
                    paid_count=r.paid_count or 0,
                    unpaid_count=(r.bill_count or 0) - (r.paid_count or 0),
                )
            )
            total_amount += total
            total_paid += paid

        overall_rate = (total_paid / total_amount * 100) if total_amount > 0 else 0

        return CollectionProgressResponse(
            items=items,
            total_amount=round(total_amount, 2),
            total_paid=round(total_paid, 2),
            total_unpaid=round(total_amount - total_paid, 2),
            overall_rate=round(overall_rate, 2),
        )

    @staticmethod
    def export_to_excel(db: Session, query: BillQuery) -> BytesIO:
        q = db.query(Bill).filter(Bill.is_deleted == False).join(Lease, Bill.lease_id == Lease.id)

        if query.keyword:
            keyword = f"%{query.keyword}%"
            q = q.filter(or_(Bill.bill_no.like(keyword), Lease.lease_no.like(keyword)))

        if query.bill_type:
            q = q.filter(Bill.bill_type == query.bill_type)

        if query.status:
            q = q.filter(Bill.status == query.status)

        bills = q.order_by(Bill.id.desc()).all()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "账单明细"

        headers = ["账单编号", "租约编号", "账单类型", "账期", "账单日期", "到期日期", "金额", "已付金额", "状态", "备注"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        for row, bill in enumerate(bills, 2):
            ws.cell(row=row, column=1, value=bill.bill_no)
            ws.cell(row=row, column=2, value=bill.lease.lease_no if bill.lease else "")
            ws.cell(row=row, column=3, value=bill.bill_type)
            ws.cell(row=row, column=4, value=bill.bill_period or "")
            ws.cell(row=row, column=5, value=str(bill.bill_date))
            ws.cell(row=row, column=6, value=str(bill.due_date))
            ws.cell(row=row, column=7, value=float(bill.amount))
            ws.cell(row=row, column=8, value=float(bill.paid_amount))
            ws.cell(row=row, column=9, value=bill.status)
            ws.cell(row=row, column=10, value=bill.remark or "")

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output


class BillPaymentService:
    @staticmethod
    def get(db: Session, payment_id: int) -> Optional[BillPayment]:
        return db.query(BillPayment).filter(BillPayment.id == payment_id, BillPayment.is_deleted == False).first()

    @staticmethod
    def list_by_bill(db: Session, bill_id: int) -> List[BillPayment]:
        return (
            db.query(BillPayment)
            .filter(BillPayment.bill_id == bill_id, BillPayment.is_deleted == False)
            .order_by(BillPayment.id.desc())
            .all()
        )

    @staticmethod
    def generate_payment_no() -> str:
        date_str = datetime.now().strftime("%Y%m%d")
        unique_id = str(uuid.uuid4().hex)[:8].upper()
        return f"SK{date_str}{unique_id}"

    @staticmethod
    def create(db: Session, data: BillPaymentCreate, created_by: Optional[int] = None) -> BillPayment:
        payment_no = BillPaymentService.generate_payment_no()
        db_obj = BillPayment(**data.model_dump(), payment_no=payment_no, created_by=created_by)
        db.add(db_obj)

        bill = db.query(Bill).filter(Bill.id == data.bill_id).first()
        if bill:
            bill.paid_amount = float(bill.paid_amount) + float(data.amount)
            if bill.paid_amount >= float(bill.amount):
                bill.status = "paid"
            else:
                bill.status = "partial"
            bill.updated_at = datetime.utcnow()

        db.commit()
        db.refresh(db_obj)
        return db_obj

    @staticmethod
    def delete(db: Session, payment_id: int, updated_by: Optional[int] = None) -> bool:
        db_obj = BillPaymentService.get(db, payment_id)
        if not db_obj:
            return False

        bill = db.query(Bill).filter(Bill.id == db_obj.bill_id).first()
        if bill:
            bill.paid_amount = max(0, float(bill.paid_amount) - float(db_obj.amount))
            if bill.paid_amount >= float(bill.amount):
                bill.status = "paid"
            elif bill.paid_amount > 0:
                bill.status = "partial"
            else:
                bill.status = "pending"
            bill.updated_at = datetime.utcnow()

        db_obj.is_deleted = True
        db_obj.updated_by = updated_by
        db_obj.updated_at = datetime.utcnow()
        db.commit()
        return True
