from fastapi import APIRouter, Depends, Query, Response
from sqlalchemy.orm import Session
from typing import Optional
from datetime import date
from app.database import get_db
from app.schemas import FilterParams
from app.services.analytics import (
    get_floor_heatmap, get_rework_trend, get_shift_comparison,
    get_work_orders_detail, get_cleaner_performance
)
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from io import BytesIO

router = APIRouter(prefix="/api", tags=["analytics"])


def build_filters(
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    floor: Optional[int] = Query(None),
    shift: Optional[str] = Query(None),
    is_vip: Optional[bool] = Query(None),
    is_late_checkout: Optional[bool] = Query(None),
    cleaner_id: Optional[int] = Query(None),
) -> FilterParams:
    return FilterParams(
        start_date=start_date,
        end_date=end_date,
        floor=floor,
        shift=shift,
        is_vip=is_vip,
        is_late_checkout=is_late_checkout,
        cleaner_id=cleaner_id,
    )


@router.get("/floor-heatmap")
def floor_heatmap(filters: FilterParams = Depends(build_filters), db: Session = Depends(get_db)):
    return get_floor_heatmap(db, filters)


@router.get("/rework-trend")
def rework_trend(filters: FilterParams = Depends(build_filters), db: Session = Depends(get_db)):
    return get_rework_trend(db, filters)


@router.get("/shift-comparison")
def shift_comparison(filters: FilterParams = Depends(build_filters), db: Session = Depends(get_db)):
    return get_shift_comparison(db, filters)


@router.get("/work-orders")
def work_orders(filters: FilterParams = Depends(build_filters), db: Session = Depends(get_db)):
    return get_work_orders_detail(db, filters)


@router.get("/cleaner-performance")
def cleaner_performance(filters: FilterParams = Depends(build_filters), db: Session = Depends(get_db)):
    return get_cleaner_performance(db, filters)


@router.get("/rooms")
def get_rooms(db: Session = Depends(get_db)):
    from app.models.models import Room
    rooms = db.query(Room).all()
    return [
        {
            "id": r.id,
            "room_number": r.room_number,
            "floor": r.floor,
            "room_type": r.room_type,
            "is_vip": r.is_vip,
            "status": r.status,
        }
        for r in rooms
    ]


@router.get("/cleaners")
def get_cleaners(db: Session = Depends(get_db)):
    from app.models.models import Cleaner
    cleaners = db.query(Cleaner).filter(Cleaner.is_active == True).all()
    return [
        {
            "id": c.id,
            "name": c.name,
            "employee_id": c.employee_id,
            "shift": c.shift,
        }
        for c in cleaners
    ]


@router.get("/weekly-report-export")
def weekly_report_export(filters: FilterParams = Depends(build_filters), db: Session = Depends(get_db)):
    wb = Workbook()

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    ws1 = wb.active
    ws1.title = "保洁员绩效"
    perf = get_cleaner_performance(db, filters)
    perf_headers = ["保洁员", "班次", "总工单", "VIP工单", "普通工单", "平均时长(min)", "VIP平均时长", "普通平均时长",
                    "返工次数", "返工率%", "VIP返工率%", "普通返工率%", "换班次数", "交接转出时长", "接手转入时长", "平均返工距查房(min)"]
    ws1.append(perf_headers)
    for col in range(1, len(perf_headers) + 1):
        cell = ws1.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    for p in perf:
        ws1.append([
            p.cleaner_name, p.shift, p.total_orders, p.vip_orders, p.normal_orders,
            p.avg_duration, p.avg_duration_vip, p.avg_duration_normal,
            p.rework_count, p.rework_rate, p.rework_rate_vip, p.rework_rate_normal,
            p.handover_count, p.handover_from_duration_total, p.handover_to_duration_total,
            p.avg_minutes_after_inspection,
        ])

    ws2 = wb.create_sheet("楼层热力图")
    heatmap = get_floor_heatmap(db, filters)
    hm_headers = ["房号", "楼层", "房型", "VIP", "延迟退房", "状态", "平均清洁时长", "返工次数"]
    ws2.append(hm_headers)
    for col in range(1, len(hm_headers) + 1):
        cell = ws2.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    for h in heatmap:
        ws2.append([
            h.room_number, h.floor, h.room_type, h.is_vip, h.is_late_checkout,
            h.status, h.avg_duration, h.rework_count,
        ])

    ws3 = wb.create_sheet("返工率趋势")
    trend = get_rework_trend(db, filters)
    trend_headers = ["日期", "VIP返工率%", "普通返工率%", "总返工率%", "VIP平均时长", "普通平均时长"]
    ws3.append(trend_headers)
    for col in range(1, len(trend_headers) + 1):
        cell = ws3.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    for t in trend:
        ws3.append([
            t.date, t.rework_rate_vip, t.rework_rate_normal, t.rework_rate_total,
            t.avg_duration_vip, t.avg_duration_normal,
        ])

    ws4 = wb.create_sheet("班次对比")
    shifts = get_shift_comparison(db, filters)
    shift_headers = ["班次", "总工单", "VIP平均时长", "普通平均时长", "VIP返工率%", "普通返工率%"]
    ws4.append(shift_headers)
    for col in range(1, len(shift_headers) + 1):
        cell = ws4.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    for s in shifts:
        ws4.append([
            s.shift, s.total_orders, s.avg_duration_vip, s.avg_duration_normal,
            s.rework_rate_vip, s.rework_rate_normal,
        ])

    ws5 = wb.create_sheet("工单明细")
    orders = get_work_orders_detail(db, filters)
    order_headers = ["工单号", "房号", "楼层", "房型", "VIP", "延迟退房", "保洁员", "班次",
                     "状态", "清洁时长", "返工次数", "返工原因", "换班记录"]
    ws5.append(order_headers)
    for col in range(1, len(order_headers) + 1):
        cell = ws5.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    for o in orders:
        rework_summary = "; ".join(
            [f"{r['reason']}(缺:{r['missing_item'] or '无'}, 距查房:{r['minutes_after_inspection']}min)" for r in o["reworks"]]
        ) if o["reworks"] else "无"
        handover_summary = "; ".join(
            [f"{h['from_cleaner_name']}->{h['to_cleaner_name']}({h['from_duration']}+{h['to_duration']}min)" for h in o["handovers"]]
        ) if o["handovers"] else "无"
        ws5.append([
            o["order_number"], o["room_number"], o["floor"], o["room_type"],
            o["is_vip"], o["is_late_checkout"], o["cleaner_name"], o["shift"],
            o["status"], o["cleaning_duration"], len(o["reworks"]),
            rework_summary, handover_summary,
        ])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"weekly_report_{filters.start_date or 'all'}_{filters.end_date or 'all'}.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
