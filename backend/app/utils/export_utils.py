from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session
from app.models import NotificationReceipt, SeatViolation, OperationHistory, User


def get_notification_info(db: Session, order_id: int, student_id: int) -> dict:
    notifications = db.query(NotificationReceipt).filter(
        NotificationReceipt.order_id == order_id
    ).order_by(NotificationReceipt.sent_at.desc()).all()
    if notifications:
        latest = notifications[0]
        return {
            "has_notification": "是",
            "latest_channel": latest.channel,
            "latest_content": latest.content or "",
            "latest_sent_at": latest.sent_at.strftime("%Y-%m-%d %H:%M:%S") if latest.sent_at else "",
            "notification_count": len(notifications),
        }
    return {
        "has_notification": "否",
        "latest_channel": "",
        "latest_content": "",
        "latest_sent_at": "",
        "notification_count": 0,
    }


def get_seat_violation_info(db: Session, student_id: int) -> dict:
    violations = db.query(SeatViolation).filter(
        SeatViolation.student_id == student_id
    ).order_by(SeatViolation.created_at.desc()).all()
    if violations:
        latest = violations[0]
        return {
            "has_violation": "是",
            "violation_count": len(violations),
            "latest_violation_type": latest.violation_type or "",
            "latest_seat_number": latest.seat_number or "",
            "latest_library_room": latest.library_room or "",
            "latest_violation_at": latest.created_at.strftime("%Y-%m-%d %H:%M:%S") if latest.created_at else "",
        }
    return {
        "has_violation": "否",
        "violation_count": 0,
        "latest_violation_type": "",
        "latest_seat_number": "",
        "latest_library_room": "",
        "latest_violation_at": "",
    }


def get_latest_change_info(db: Session, order_id: int) -> dict:
    histories = db.query(OperationHistory).filter(
        OperationHistory.order_id == order_id
    ).order_by(OperationHistory.created_at.desc()).all()
    if histories:
        latest = histories[0]
        operator = db.query(User).filter(User.id == latest.operator_id).first()
        return {
            "has_change": "是",
            "change_count": len(histories),
            "latest_field": latest.field_name,
            "latest_old_value": latest.old_value or "",
            "latest_new_value": latest.new_value or "",
            "latest_operator": operator.real_name or operator.username if operator else "",
            "latest_change_at": latest.created_at.strftime("%Y-%m-%d %H:%M:%S") if latest.created_at else "",
        }
    return {
        "has_change": "否",
        "change_count": 0,
        "latest_field": "",
        "latest_old_value": "",
        "latest_new_value": "",
        "latest_operator": "",
        "latest_change_at": "",
    }


def generate_excel(queryset, columns: list, filter_params: dict, operator_name: str, file_path: str, db: Session):
    wb = Workbook()

    ws_data = wb.active
    ws_data.title = "报修数据"

    category_map = {
        "plumbing": "水管",
        "electrical": "电路",
        "furniture": "家具",
        "door_window": "门窗",
        "other": "其他",
    }
    status_map = {
        "pending": "待审核",
        "in_review": "审核中",
        "approved": "已通过",
        "rejected": "已驳回",
        "in_progress": "处理中",
        "completed": "已完成",
        "closed": "已关闭",
    }
    urgency_map = {
        "low": "低",
        "medium": "中",
        "high": "高",
    }
    channel_map = {
        "sms": "短信",
        "email": "邮件",
        "in_app": "站内信",
    }

    full_headers = [
        "工单号", "标题", "描述", "类别", "状态", "紧急程度",
        "宿舍房间", "位置", "提交人", "学号", "联系方式",
        "创建时间", "更新时间",
        "消息触达", "最新通知方式", "最新通知内容", "最新通知时间", "通知次数",
        "座位违约", "违约次数", "最新违约类型", "最新座位号", "最新自习室", "最新违约时间",
        "最近变更", "变更次数", "最新变更字段", "旧值", "新值", "操作人", "变更时间",
    ]

    header_fill = PatternFill(start_color="2080f0", end_color="2080f0", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for col_idx, header in enumerate(full_headers, 1):
        cell = ws_data.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font

    for row_idx, item in enumerate(queryset, 2):
        student = db.query(User).filter(User.id == item.student_id).first()
        notif_info = get_notification_info(db, item.id, item.student_id)
        violation_info = get_seat_violation_info(db, item.student_id)
        change_info = get_latest_change_info(db, item.id)

        row_data = [
            item.id,
            item.title,
            item.description or "",
            category_map.get(item.category, item.category),
            status_map.get(item.status, item.status),
            urgency_map.get(item.urgency, item.urgency),
            item.dorm_room or "",
            item.location or "",
            student.real_name if student else "",
            student.student_id if student else "",
            student.phone if student else "",
            item.created_at.strftime("%Y-%m-%d %H:%M:%S") if item.created_at else "",
            item.updated_at.strftime("%Y-%m-%d %H:%M:%S") if item.updated_at else "",
            notif_info["has_notification"],
            channel_map.get(notif_info["latest_channel"], notif_info["latest_channel"]),
            notif_info["latest_content"],
            notif_info["latest_sent_at"],
            notif_info["notification_count"],
            violation_info["has_violation"],
            violation_info["violation_count"],
            violation_info["latest_violation_type"],
            violation_info["latest_seat_number"],
            violation_info["latest_library_room"],
            violation_info["latest_violation_at"],
            change_info["has_change"],
            change_info["change_count"],
            change_info["latest_field"],
            change_info["latest_old_value"],
            change_info["latest_new_value"],
            change_info["latest_operator"],
            change_info["latest_change_at"],
        ]

        for col_idx, value in enumerate(row_data, 1):
            ws_data.cell(row=row_idx, column=col_idx, value=value)

    for col_idx in range(1, len(full_headers) + 1):
        ws_data.column_dimensions[chr(64 + col_idx) if col_idx <= 26 else "A" + chr(64 + col_idx - 26)].width = 18

    ws_meta = wb.create_sheet(title="元数据")
    ws_meta.append(["字段", "值"])
    ws_meta["A1"].font = Font(bold=True)
    ws_meta["B1"].font = Font(bold=True)
    ws_meta.append(["生成时间", datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S")])
    ws_meta.append(["操作人", operator_name])
    ws_meta.append(["总记录数", len(queryset)])
    ws_meta.append(["导出类型", "报修数据"])

    ws_filter = wb.create_sheet(title="筛选口径")
    ws_filter.append(["字段", "值"])
    ws_filter["A1"].font = Font(bold=True)
    ws_filter["B1"].font = Font(bold=True)
    if filter_params:
        for key, value in filter_params.items():
            key_display = {
                "status": "状态",
                "category": "类别",
                "student_id": "学生ID",
                "start_date": "开始日期",
                "end_date": "结束日期",
            }.get(key, key)
            ws_filter.append([key_display, str(value)])
    else:
        ws_filter.append(["未应用筛选条件", ""])

    ws_notifications = wb.create_sheet(title="通知回执明细")
    notif_headers = ["工单号", "通知方式", "通知内容", "是否已读", "发送时间"]
    for col_idx, header in enumerate(notif_headers, 1):
        cell = ws_notifications.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font

    all_notifications = db.query(NotificationReceipt).filter(
        NotificationReceipt.order_id.in_([item.id for item in queryset])
    ).all()
    for row_idx, notif in enumerate(all_notifications, 2):
        ws_notifications.cell(row=row_idx, column=1, value=notif.order_id)
        ws_notifications.cell(row=row_idx, column=2, value=channel_map.get(notif.channel, notif.channel))
        ws_notifications.cell(row=row_idx, column=3, value=notif.content or "")
        ws_notifications.cell(row=row_idx, column=4, value="是" if notif.is_read else "否")
        ws_notifications.cell(row=row_idx, column=5, value=notif.sent_at.strftime("%Y-%m-%d %H:%M:%S") if notif.sent_at else "")

    ws_violations = wb.create_sheet(title="座位违约明细")
    viol_headers = ["学号", "姓名", "座位号", "自习室", "违约类型", "违约时间"]
    for col_idx, header in enumerate(viol_headers, 1):
        cell = ws_violations.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font

    student_ids = list(set([item.student_id for item in queryset]))
    all_violations = db.query(SeatViolation).filter(
        SeatViolation.student_id.in_(student_ids)
    ).all()
    for row_idx, viol in enumerate(all_violations, 2):
        s = db.query(User).filter(User.id == viol.student_id).first()
        ws_violations.cell(row=row_idx, column=1, value=s.student_id if s else "")
        ws_violations.cell(row=row_idx, column=2, value=s.real_name if s else "")
        ws_violations.cell(row=row_idx, column=3, value=viol.seat_number or "")
        ws_violations.cell(row=row_idx, column=4, value=viol.library_room or "")
        ws_violations.cell(row=row_idx, column=5, value=viol.violation_type or "")
        ws_violations.cell(row=row_idx, column=6, value=viol.created_at.strftime("%Y-%m-%d %H:%M:%S") if viol.created_at else "")

    ws_changes = wb.create_sheet(title="变更历史明细")
    change_headers = ["工单号", "变更字段", "旧值", "新值", "操作人", "变更时间"]
    for col_idx, header in enumerate(change_headers, 1):
        cell = ws_changes.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font

    all_changes = db.query(OperationHistory).filter(
        OperationHistory.order_id.in_([item.id for item in queryset])
    ).all()
    for row_idx, change in enumerate(all_changes, 2):
        op = db.query(User).filter(User.id == change.operator_id).first()
        ws_changes.cell(row=row_idx, column=1, value=change.order_id)
        ws_changes.cell(row=row_idx, column=2, value=change.field_name)
        ws_changes.cell(row=row_idx, column=3, value=change.old_value or "")
        ws_changes.cell(row=row_idx, column=4, value=change.new_value or "")
        ws_changes.cell(row=row_idx, column=5, value=op.real_name if op else "")
        ws_changes.cell(row=row_idx, column=6, value=change.created_at.strftime("%Y-%m-%d %H:%M:%S") if change.created_at else "")

    wb.save(file_path)
    return file_path
