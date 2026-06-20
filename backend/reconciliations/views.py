import csv
import io
from decimal import Decimal

from openpyxl import load_workbook
from rest_framework import status, viewsets
from rest_framework.decorators import action
from rest_framework.permissions import BasePermission, IsAuthenticated
from rest_framework.response import Response

from .models import Difference, Reconciliation
from .serializers import (
    DifferenceSerializer,
    ReconciliationSerializer,
    ReconciliationUploadSerializer,
)


class IsProjectManagerOrAdmin(BasePermission):
    def has_permission(self, request, view):
        if not request.user.is_authenticated:
            return False
        return request.user.role in ('admin', 'project_manager')


class IsFinanceOrAbove(BasePermission):
    def has_permission(self, request, view):
        if not request.user.is_authenticated:
            return False
        return request.user.role in ('admin', 'project_manager', 'finance')


class ReconciliationViewSet(viewsets.ModelViewSet):
    serializer_class = ReconciliationSerializer
    queryset = Reconciliation.objects.all()

    def get_permissions(self):
        if self.action in ('create', 'confirm', 'reject'):
            return [IsAuthenticated(), IsProjectManagerOrAdmin()]
        return [IsAuthenticated(), IsFinanceOrAbove()]

    def create(self, request, *args, **kwargs):
        upload_serializer = ReconciliationUploadSerializer(data=request.data)
        upload_serializer.is_valid(raise_exception=True)

        uploaded_file = upload_serializer.validated_data['file']
        project_name = upload_serializer.validated_data['project_name']
        client_name = upload_serializer.validated_data['client_name']

        reconciliation = Reconciliation.objects.create(
            project_name=project_name,
            client_name=client_name,
            uploaded_by=request.user,
            file=uploaded_file,
        )

        differences = self._parse_file(uploaded_file, reconciliation)

        total_amount = Decimal('0')
        matched_amount = Decimal('0')
        for diff in differences:
            try:
                val = Decimal(diff.system_value)
                total_amount += val
                if diff.item_type == Difference.ItemType.AMOUNT:
                    matched_amount += Decimal(diff.uploaded_value)
            except Exception:
                pass

        reconciliation.total_amount = total_amount
        reconciliation.matched_amount = matched_amount
        reconciliation.difference_amount = total_amount - matched_amount
        if differences:
            reconciliation.status = Reconciliation.Status.COMPARED
        reconciliation.save()

        output_serializer = ReconciliationSerializer(reconciliation)
        return Response(output_serializer.data, status=status.HTTP_201_CREATED)

    def _parse_file(self, uploaded_file, reconciliation):
        filename = uploaded_file.name.lower()
        differences = []

        if filename.endswith(('.xlsx', '.xlsm', '.xltx', '.xltm')):
            differences = self._parse_excel(uploaded_file, reconciliation)
        elif filename.endswith('.csv'):
            differences = self._parse_csv(uploaded_file, reconciliation)

        return differences

    def _parse_excel(self, uploaded_file, reconciliation):
        differences = []
        wb = load_workbook(filename=uploaded_file, read_only=True)
        ws = wb.active

        headers = None
        for row in ws.iter_rows(values_only=True):
            if headers is None:
                headers = [str(cell).strip().lower() if cell else '' for cell in row]
                continue

            row_data = {}
            for i, cell in enumerate(row):
                if i < len(headers):
                    row_data[headers[i]] = str(cell) if cell is not None else ''

            diff = self._row_to_difference(row_data, reconciliation)
            if diff:
                differences.append(diff)

        wb.close()
        Difference.objects.bulk_create(differences)
        return differences

    def _parse_csv(self, uploaded_file, reconciliation):
        differences = []
        decoded = uploaded_file.read().decode('utf-8')
        reader = csv.DictReader(io.StringIO(decoded))

        for row in reader:
            row_data = {k.strip().lower(): str(v) if v else '' for k, v in row.items()}
            diff = self._row_to_difference(row_data, reconciliation)
            if diff:
                differences.append(diff)

        Difference.objects.bulk_create(differences)
        return differences

    def _row_to_difference(self, row_data, reconciliation):
        item_type = row_data.get('item_type', row_data.get('type', 'amount'))
        system_value = row_data.get('system_value', row_data.get('system', ''))
        uploaded_value = row_data.get('uploaded_value', row_data.get('uploaded', ''))

        if not system_value and not uploaded_value:
            return None

        valid_types = [t[0] for t in Difference.ItemType.choices]
        if item_type not in valid_types:
            item_type = Difference.ItemType.AMOUNT

        return Difference(
            reconciliation=reconciliation,
            item_type=item_type,
            system_value=system_value,
            uploaded_value=uploaded_value,
        )

    @action(detail=True, methods=['post'])
    def confirm(self, request, pk=None):
        reconciliation = self.get_object()

        if reconciliation.status != Reconciliation.Status.COMPARED:
            return Response(
                {'detail': 'Only compared reconciliations can be confirmed.'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        from django.utils import timezone

        differences = reconciliation.differences.all()
        differences.update(
            is_confirmed=True,
            confirmed_by=request.user,
            confirmed_at=timezone.now(),
        )

        reconciliation.status = Reconciliation.Status.CONFIRMED
        reconciliation.save()

        from reminders.tasks import generate_reminders
        generate_reminders.delay(reconciliation.id)

        return Response(ReconciliationSerializer(reconciliation).data)

    @action(detail=True, methods=['post'])
    def reject(self, request, pk=None):
        reconciliation = self.get_object()

        if reconciliation.status not in (Reconciliation.Status.COMPARED, Reconciliation.Status.PENDING):
            return Response(
                {'detail': 'Only compared or pending reconciliations can be rejected.'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        reconciliation.status = Reconciliation.Status.REJECTED
        reconciliation.save()

        return Response(ReconciliationSerializer(reconciliation).data)

    @action(detail=True, methods=['get'])
    def differences(self, request, pk=None):
        reconciliation = self.get_object()
        diffs = reconciliation.differences.all()
        serializer = DifferenceSerializer(diffs, many=True)
        return Response(serializer.data)
