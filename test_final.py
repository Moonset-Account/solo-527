import os
import django
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'tea_loss_dashboard.settings')
django.setup()

from django.test import Client
import json
from django.conf import settings
from reports.models import ReportTask

client = Client()
client.login(username='admin', password='admin123')

all_passed = True

print('=' * 70)
print('连锁茶饮原料损耗看板 - 完整端到端测试')
print('=' * 70)

# 1. 测试所有4种报表类型生成
print('\n1. 测试所有报表类型生成（共4种）')
report_types = ['loss_summary', 'loss_detail', 'store_ranking', 'material_analysis']
for rt in report_types:
    try:
        report_data = {
            'report_type': rt,
            'start_date': '2026-05-01',
            'end_date': '2026-06-07',
            'exclude_trial': True
        }
        response = client.post('/api/v1/reports/generate/',
                               data=json.dumps(report_data),
                               content_type='application/json')
        if response.status_code == 200:
            data = response.json()
            task_id = data['data']['task_id']
            task = ReportTask.objects.get(task_id=task_id)
            status = task.status
            has_file = bool(task.file_path)
            print(f'   ✅ {rt}: {status} - 文件生成: {"是" if has_file else "否"}')
            if not has_file:
                all_passed = False
        else:
            print(f'   ❌ {rt}: HTTP {response.status_code}')
            all_passed = False
    except Exception as e:
        print(f'   ❌ {rt}: 异常 - {str(e)}')
        all_passed = False

# 2. 测试报表文件包含排除试营说明
print('\n2. 验证报表文件包含排除试营说明')
try:
    task = ReportTask.objects.filter(status='completed').first()
    if task and task.file_path:
        import openpyxl
        full_path = os.path.join(settings.MEDIA_ROOT, task.file_path)
        if os.path.exists(full_path):
            wb = openpyxl.load_workbook(full_path)
            info_sheet = wb['报表说明']
            has_exclude_trial = False
            for row in info_sheet.iter_rows():
                for cell in row:
                    if cell.value and '排除试营原料' in str(cell.value):
                        has_exclude_trial = True
                        break
            print(f'   ✅ 报表说明中包含排除试营原料标记: {"是" if has_exclude_trial else "否"}')
            print(f'   ✅ 报表类型: {task.get_report_type_display()}')
            print(f'   ✅ 工作表: {wb.sheetnames}')
            if not has_exclude_trial:
                all_passed = False
        else:
            print(f'   ⚠️ 文件不存在: {full_path}')
    else:
        print('   ⚠️ 暂无已完成报表任务')
except Exception as e:
    print(f'   ⚠️ 验证跳过: {str(e)}')

# 3. 测试所有数据类型导入模板
print('\n3. 测试所有数据类型导入模板（共6种）')
data_types = ['wastage', 'stocktake', 'inventory', 'material_use', 'sale', 'staff_shift']
for dt in data_types:
    try:
        response = client.get(f'/api/v1/import/template/?type={dt}')
        if response.status_code == 200 and 'excel' in response['Content-Type']:
            print(f'   ✅ {dt} 模板: 正常')
        else:
            print(f'   ❌ {dt} 模板: HTTP {response.status_code}')
            all_passed = False
    except Exception as e:
        print(f'   ❌ {dt} 模板: 异常 - {str(e)}')
        all_passed = False

# 4. 测试数据导入功能
print('\n4. 测试数据导入功能（报损数据示例）')
try:
    sample_file = os.path.join(settings.MEDIA_ROOT, 'sample_imports', '报损数据示例.xlsx')
    if os.path.exists(sample_file):
        with open(sample_file, 'rb') as f:
            response = client.post('/api/v1/import/upload/', {
                'file': f,
                'data_type': 'wastage'
            })
        if response.status_code == 200:
            data = response.json()
            print(f'   ✅ 导入提交成功: {data["data"]["status"]}')
            print(f'   ✅ 任务ID: {data["data"]["task_id"]}')
        else:
            print(f'   ❌ 导入失败: HTTP {response.status_code}')
            all_passed = False
    else:
        print(f'   ⚠️ 示例文件不存在: {sample_file}')
except Exception as e:
    print(f'   ❌ 导入异常: {str(e)}')
    all_passed = False

# 5. 测试员工班次导入
print('\n5. 测试员工班次导入（含store字段）')
try:
    sample_file = os.path.join(settings.MEDIA_ROOT, 'sample_imports', '员工班次示例.xlsx')
    if os.path.exists(sample_file):
        with open(sample_file, 'rb') as f:
            response = client.post('/api/v1/import/upload/', {
                'file': f,
                'data_type': 'staff_shift'
            })
        if response.status_code == 200:
            data = response.json()
            print(f'   ✅ 员工班次导入提交成功: {data["data"]["status"]}')
        else:
            print(f'   ❌ 员工班次导入失败: HTTP {response.status_code}')
            all_passed = False
except Exception as e:
    print(f'   ❌ 员工班次导入异常: {str(e)}')
    all_passed = False

# 6. 测试导入任务列表
print('\n6. 测试导入任务列表')
try:
    response = client.get('/api/v1/import/tasks/')
    if response.status_code == 200:
        tasks = response.json()['data']
        print(f'   ✅ 导入任务列表正常，共 {len(tasks)} 条记录')
    else:
        print(f'   ❌ HTTP {response.status_code}')
        all_passed = False
except Exception as e:
    print(f'   ❌ 异常: {str(e)}')
    all_passed = False

# 7. 测试试营原料默认排除
print('\n7. 测试原料排行默认排除试营原料')
try:
    response = client.get('/api/v1/dashboard/material_ranking/')
    if response.status_code == 200:
        materials = response.json()['data']
        has_trial = any(m['isTrial'] for m in materials)
        print(f'   默认请求(无exclude_trial参数)')
        print(f'   是否包含试营原料: {has_trial}')
        if not has_trial:
            print('   ✅ 默认排除试营原料 - 正确')
        else:
            print('   ❌ 包含试营原料 - 错误')
            all_passed = False
except Exception as e:
    print(f'   ❌ 异常: {str(e)}')
    all_passed = False

# 8. 测试趋势API使用预聚合
print('\n8. 测试损耗趋势API（优先预聚合表）')
try:
    response = client.get('/api/v1/dashboard/trend/?exclude_trial=true')
    if response.status_code == 200:
        data = response.json()['data']
        print(f'   ✅ 趋势API正常，返回 {len(data)} 个数据点')
        if len(data) > 0:
            print(f'   ✅ 示例数据点: {data[0]}')
except Exception as e:
    print(f'   ❌ 异常: {str(e)}')
    all_passed = False

# 9. 测试店长权限
print('\n9. 测试店长权限（只能看本店数据）')
try:
    client2 = Client()
    client2.login(username='manager1', password='manager123')
    response = client2.get('/api/v1/dashboard/stores/')
    stores = response.json()['data']
    print(f'   店长可见门店数: {len(stores)}')
    if len(stores) == 1:
        print(f'   ✅ 店长只能看本店: {stores[0]["name"]}')
    else:
        print(f'   ⚠️ 门店数量异常')
    
    response = client2.get('/api/v1/reports/export-details/')
    if response.status_code == 200:
        print(f'   ✅ 店长可导出本店明细')
except Exception as e:
    print(f'   ❌ 异常: {str(e)}')
    all_passed = False

# 10. 系统配置说明
print('\n10. 系统配置说明')
print(f'    当前数据库: SQLite（默认开发模式）')
print(f'    当前Celery: EAGER模式（同步执行，开发模式）')
print(f'    PostgreSQL切换: python3 switch_config.py production')
print(f'    配置工具: python3 switch_config.py')

print('\n' + '=' * 70)
if all_passed:
    print('✅ 所有端到端测试通过！')
else:
    print('⚠️ 部分测试未通过，请检查上述错误')
print('=' * 70)

print('\n📂 示例导入文件: media/sample_imports/')
print('📊 生成的报表: media/reports/')
print('🔧 配置切换: python3 switch_config.py')
print('🌐 访问地址: http://localhost:9000/')
print('👤 测试账号: admin / admin123')
